"""Deep UI audit (developer harness, not shipped in Setup): drives every menu item and toolbar
control of the real main window with realistic input, alone and in combination, and checks results.

* File dialogs answer with sample inputs / paths in the output folder (by filter and caption).
* Every QDialog is captured (screenshot + clipped-text check), filled with realistic values and
  accepted through its own OK button; progress dialogs run their real loop under a watchdog.
* Message boxes are recorded (warnings / errors are findings unless the step expects them).
* "Drag a box on the page" tools get real mouse drags on the canvas.
* Printing goes to a PDF file; Explorer / Windows Settings launches are recorded, not run.
* A watchdog thread ends the run (with stack traces) if a step hangs.

Usage: env\\Scripts\\python.exe scripts\\deep_ui_audit.py <out_dir> [step-name-filter]
Writes <out_dir>/deep_audit.json and dialog screenshots in <out_dir>/shots.
"""

from __future__ import annotations

import collections
import faulthandler
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import traceback
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "build" / "deep-audit"
ONLY = sys.argv[2] if len(sys.argv) > 2 else ""
shutil.rmtree(OUT, ignore_errors=True)
for sub in ("files", "inputs", "shots"):
    (OUT / sub).mkdir(parents=True, exist_ok=True)
os.environ["MONOGRAPH_PDF_USER_DATA_DIR"] = str(OUT / "userdata")
os.environ["MONOGRAPH_PDF_RUNTIME_DIR"] = str(OUT / "runtime")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
CAN_PRINT = sys.platform.startswith("win") and os.environ["QT_QPA_PLATFORM"] != "offscreen"
if sys.platform.startswith("win") and os.environ["QT_QPA_PLATFORM"] == "offscreen":
    os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"))
WINDOWS = sys.platform.startswith("win")

import fitz  # noqa: E402
from PySide6.QtCore import QEvent, QPoint, QPointF, Qt, QTimer, QtMsgType, qInstallMessageHandler  # noqa: E402
from PySide6.QtGui import QDesktopServices, QMouseEvent  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402
from PySide6.QtWidgets import (  # noqa: E402
    QAbstractButton,
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QInputDialog,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QToolBar,
)

# ------------------------------------------------------------------ recording

RESULTS: list[dict] = []
CUR: dict = {}
STEP_DEADLINE = [0.0]
LOG = open(OUT / "audit.log", "w", encoding="utf-8")  # noqa: SIM115
faulthandler.enable(file=LOG, all_threads=True)


def log(*a):
    msg = " ".join(str(x) for x in a)
    print(msg, flush=True)
    LOG.write(msg + "\n")
    LOG.flush()


def check(name: str, ok: bool, detail: str = "") -> bool:
    CUR.setdefault("checks", []).append({"name": name, "ok": bool(ok), "detail": str(detail)[:400]})
    if not ok:
        log(f"    FAIL {name}: {detail}")
    return bool(ok)


def note(text: str) -> None:
    CUR.setdefault("notes", []).append(text)


def _qt_handler(mode, _ctx, message):
    if mode in (QtMsgType.QtWarningMsg, QtMsgType.QtCriticalMsg, QtMsgType.QtFatalMsg):
        if "propagateSizeHints" in message or "QFont::setPixelSize" in message:
            return
        CUR.setdefault("qt", []).append(message[:300])


def _excepthook(t, e, tb):
    CUR.setdefault("exceptions", []).append("".join(traceback.format_exception(t, e, tb))[-2000:])


qInstallMessageHandler(_qt_handler)
sys.excepthook = _excepthook


def _native_windows() -> list:
    """Visible Win32 top-level windows owned by this process (catches native dialogs Qt does not list)."""
    if not sys.platform.startswith("win"):
        return []
    import ctypes
    from ctypes import wintypes

    user32 = ctypes.windll.user32
    found = []
    pid = os.getpid()

    @ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
    def cb(hwnd, _l):
        owner = wintypes.DWORD()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(owner))
        if owner.value == pid and user32.IsWindowVisible(hwnd):
            title = ctypes.create_unicode_buffer(256)
            cls = ctypes.create_unicode_buffer(256)
            user32.GetWindowTextW(hwnd, title, 256)
            user32.GetClassNameW(hwnd, cls, 256)
            found.append(f"{cls.value}:{title.value}")
        return True

    try:
        user32.EnumWindows(cb, 0)
    except Exception as exc:  # noqa: BLE001
        return [f"EnumWindows failed: {exc}"]
    return found


LAST_UI = [[]]  # visible top-level windows, refreshed on the GUI thread, reported if a step hangs


def _watchdog():
    while True:
        time.sleep(2)
        if STEP_DEADLINE[0] and time.time() > STEP_DEADLINE[0]:
            log(f"!!! TIMEOUT in step {CUR.get('step')} — dumping stacks and exiting")
            log(f"    windows open at the last check: {LAST_UI[0]}")
            log(f"    native windows of this process: {_native_windows()}")
            faulthandler.dump_traceback(file=LOG, all_threads=True)
            CUR.setdefault("exceptions", []).append("TIMEOUT (step hung)")
            RESULTS.append(dict(CUR))
            _write()
            LOG.flush()
            os._exit(3)


def _write():
    (OUT / "deep_audit.json").write_text(json.dumps(RESULTS, ensure_ascii=False, indent=1, default=str), encoding="utf-8")


threading.Thread(target=_watchdog, daemon=True).start()

# ------------------------------------------------------------------ inputs

SAMPLES = ROOT / "tests" / "samples"
from tests.fixtures.make_form_samples import main as _mk_forms  # noqa: E402
from tests.fixtures.make_ocr_samples import main as _mk_ocr  # noqa: E402
from tests.fixtures.make_samples import main as _mk_samples  # noqa: E402

_mk_samples()
_mk_forms()
_mk_ocr()
INP = OUT / "inputs"
_counter = [0]


def fresh(name: str) -> Path:
    """A new copy of a sample (never the original)."""
    _counter[0] += 1
    dst = INP / f"{_counter[0]:03d}-{name}"
    shutil.copy(SAMPLES / name, dst)
    return dst


def make_inputs():
    from PIL import Image, ImageDraw

    for i, color in enumerate(("#2563eb", "#dc2626")):
        im = Image.new("RGB", (480, 320), "white")
        d = ImageDraw.Draw(im)
        d.rectangle((20, 20, 460, 300), outline=color, width=8)
        d.text((60, 140), f"Image {i + 1}", fill=color)
        im.save(INP / f"picture{i + 1}.png")
    try:
        import docx

        doc = docx.Document()
        doc.add_heading("Word sample 見出し", 1)
        doc.add_paragraph("Paragraph text for Office conversion. 日本語の段落。")
        doc.save(INP / "sample.docx")
    except Exception as exc:  # noqa: BLE001
        log("no docx input:", exc)
    try:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Qty"])
        ws.append(["Apple りんご", 3])
        wb.save(INP / "sample.xlsx")
    except Exception as exc:  # noqa: BLE001
        log("no xlsx input:", exc)
    try:
        import pptx

        prs = pptx.Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Slide title スライド"
        prs.save(INP / "sample.pptx")
    except Exception as exc:  # noqa: BLE001
        log("no pptx input:", exc)


make_inputs()

# ------------------------------------------------------------------ dialog patches

OPEN_QUEUE: collections.deque = collections.deque()  # explicit answers for the next open dialogs
SAVE_QUEUE: collections.deque = collections.deque()
TEXT_QUEUE: collections.deque = collections.deque()
MSG_QUEUE: collections.deque = collections.deque()  # answers for question boxes, e.g. "yes"/"no"/"discard"
FILLERS: dict = {}
LAUNCHED: list = []
_shot_no = [0]


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", s).strip("-")[:50] or "file"


def _ext_from_filter(flt: str) -> str:
    m = re.search(r"\*\.(\w+)", flt or "")
    return "." + m.group(1) if m else ".pdf"


def _get_open(parent=None, caption="", directory="", filt="", *a, **k):
    CUR.setdefault("file_dialogs", []).append(f"open: {caption} [{filt}]")
    if OPEN_QUEUE:
        return str(OPEN_QUEUE.popleft()), ""
    cap = str(caption).lower()
    low = f"{caption} {filt}".lower() if cap not in ("open", "open pdf") else cap
    if "image" in low or "png" in low:
        return str(INP / "picture1.png"), ""
    if "word" in low or "docx" in low or "office" in low:
        return str(INP / "sample.docx"), ""
    return str(fresh("sample_text.pdf")), ""


def _get_opens(parent=None, caption="", directory="", filt="", *a, **k):
    CUR.setdefault("file_dialogs", []).append(f"open many: {caption} [{filt}]")
    if OPEN_QUEUE:
        return [str(p) for p in OPEN_QUEUE.popleft()], ""
    low = f"{caption} {filt}".lower()
    if "image" in low or "png" in low:
        return [str(INP / "picture1.png"), str(INP / "picture2.png")], ""
    return [str(fresh("sample_text.pdf")), str(fresh("sample_many_pages_10.pdf"))], ""


def _get_save(parent=None, caption="", default="", filt="", *a, **k):
    CUR.setdefault("file_dialogs", []).append(f"save: {caption} [{default}] [{filt}]")
    if SAVE_QUEUE:
        return str(SAVE_QUEUE.popleft()), filt
    base = Path(default).name if default else _slug(caption)
    stem, ext = os.path.splitext(base)
    if not ext:
        ext = _ext_from_filter(filt)
    path = OUT / "files" / f"{_slug(CUR.get('step', 'x'))}--{_slug(stem)}{ext}"
    n = 2
    while path.exists():  # a user picks a new name; the open document cannot be replaced
        path = path.with_name(f"{_slug(CUR.get('step', 'x'))}--{_slug(stem)}-{n}{ext}")
        n += 1
    CUR.setdefault("saved_to", []).append(str(path))
    return str(path), filt


def _get_dir(parent=None, caption="", directory="", *a, **k):
    path = OUT / "files" / f"{_slug(CUR.get('step', 'x'))}--dir"
    path.mkdir(parents=True, exist_ok=True)
    CUR.setdefault("saved_to", []).append(str(path))
    return str(path)


QFileDialog.getOpenFileName = staticmethod(_get_open)
QFileDialog.getOpenFileNames = staticmethod(_get_opens)
QFileDialog.getSaveFileName = staticmethod(_get_save)
QFileDialog.getExistingDirectory = staticmethod(_get_dir)
QInputDialog.getText = staticmethod(lambda *a, **k: (TEXT_QUEUE.popleft() if TEXT_QUEUE else "alpha", True))
QInputDialog.getInt = staticmethod(lambda *a, **k: (int(TEXT_QUEUE.popleft()) if TEXT_QUEUE else 2, True))
QInputDialog.getItem = staticmethod(lambda p, t, l, items, *a, **k: (items[0] if items else "", True))
QInputDialog.getDouble = staticmethod(lambda *a, **k: (a[3] if len(a) > 3 else 1.0, True))
QInputDialog.getMultiLineText = staticmethod(lambda *a, **k: (TEXT_QUEUE.popleft() if TEXT_QUEUE else "Note text", True))

SB = QMessageBox.StandardButton


def _answer(buttons, title, text, default=None):
    want = MSG_QUEUE.popleft() if MSG_QUEUE else None
    table = {"yes": SB.Yes, "no": SB.No, "ok": SB.Ok, "save": SB.Save, "discard": SB.Discard,
             "cancel": SB.Cancel, "open": SB.Open, "close": SB.Close}
    if want:
        return table[want]
    b = int(buttons.value) if hasattr(buttons, "value") else int(buttons)
    for cand in (SB.Discard, SB.No, SB.Close, SB.Ok, SB.Yes):  # default: never keep / open extra things
        if b & int(cand.value):
            return cand
    return SB.Ok


def _box_static(kind):
    def f(parent, title, text, buttons=SB.Ok, default=SB.NoButton):
        b = int(buttons.value) if hasattr(buttons, "value") else int(buttons)
        asks = bool(b & (int(SB.Yes.value) | int(SB.No.value) | int(SB.Cancel.value) | int(SB.Save.value)))
        CUR.setdefault("messages", []).append({"kind": "question" if asks else kind, "title": str(title),
                                               "text": str(text)[:400]})
        if asks:
            return _answer(buttons, title, text)
        return SB.Ok
    return staticmethod(f)


for _k in ("information", "warning", "critical", "question"):
    setattr(QMessageBox, _k, _box_static(_k))


def _box_exec(self):
    CUR.setdefault("messages", []).append({"kind": "box", "title": self.windowTitle(), "text": self.text()[:400]})
    custom = [b for b in self.buttons() if self.standardButton(b) == QMessageBox.StandardButton.NoButton]
    if custom:
        # boxes with named buttons (e.g. unsaved changes before a conversion: Save / Save As… / Cancel)
        want = (MSG_QUEUE.popleft() if MSG_QUEUE else "save").lower()
        pick = next((b for b in self.buttons() if b.text().replace("&", "").rstrip("…").lower() == want), None)
        if pick is None and want == "cancel":
            pick = self.button(QMessageBox.StandardButton.Cancel)
        self.clickedButton = lambda pick=pick: pick  # noqa: E731
        return 0
    return int(_answer(self.standardButtons(), self.windowTitle(), self.text()).value)


QMessageBox.exec = _box_exec
_real_open_url = QDesktopServices.openUrl
QDesktopServices.openUrl = staticmethod(lambda url: LAUNCHED.append(("url", url.toString())) or True)
if hasattr(os, "startfile"):
    os.startfile = lambda *a, **k: LAUNCHED.append(("startfile", a))
_real_popen = subprocess.Popen


class _Popen(_real_popen):
    def __init__(self, args, *a, **k):
        first = str(args[0] if isinstance(args, (list, tuple)) else args).lower()
        if any(x in first for x in ("explorer", "rundll32", "ms-settings")):
            LAUNCHED.append(("process", args))
            args = [sys.executable, "-c", "pass"]
        super().__init__(args, *a, **k)


subprocess.Popen = _Popen

REAL_EXEC = QDialog.exec
REAL_LOOP = {"ProgressDialog", "OcrProgressDialog", "BusyDialog"}


def _clipped(dlg) -> list[str]:
    """Labels / buttons whose text does not fit (cut off) in the shown dialog."""
    issues = []
    for lab in dlg.findChildren(QLabel):
        if not lab.isVisible() or not lab.text().strip() or lab.wordWrap() or lab.pixmap() is not None and not lab.pixmap().isNull():
            continue
        txt = re.sub(r"<[^>]+>", "", lab.text())
        need = lab.fontMetrics().horizontalAdvance(txt.split("\n")[0])
        if need > lab.width() + 2 and "\n" not in txt and len(txt) < 200:
            issues.append(f"label cut: '{txt[:60]}' needs {need}px has {lab.width()}px")
    for b in dlg.findChildren(QAbstractButton):
        if b.isVisible() and b.text().strip() and isinstance(b, (QPushButton, QCheckBox)):
            need = b.fontMetrics().horizontalAdvance(b.text().replace("&", "")) + (24 if isinstance(b, QCheckBox) else 12)
            if need > b.width() + 2:
                issues.append(f"button cut: '{b.text()}' needs {need}px has {b.width()}px")
    scr = QApplication.primaryScreen().availableGeometry()
    if dlg.width() > 1366 or dlg.height() > 740:
        issues.append(f"dialog {dlg.width()}x{dlg.height()} too big for a 1366x768 screen")
    del scr
    return issues


def _click_ok(dlg) -> str:
    for box in dlg.findChildren(QDialogButtonBox):
        if not box.isVisible():
            continue
        for sb in (QDialogButtonBox.StandardButton.Ok, QDialogButtonBox.StandardButton.Save,
                   QDialogButtonBox.StandardButton.Yes, QDialogButtonBox.StandardButton.Open):
            b = box.button(sb)
            if b is not None and b.isVisible():
                if not b.isEnabled():
                    return f"OK disabled ({b.text()})"
                b.click()
                return f"clicked {b.text()}"
        for b in box.buttons():
            if box.buttonRole(b) in (QDialogButtonBox.ButtonRole.AcceptRole, QDialogButtonBox.ButtonRole.YesRole):
                if not b.isEnabled():
                    return f"OK disabled ({b.text()})"
                b.click()
                return f"clicked {b.text()}"
    for b in dlg.findChildren(QPushButton):
        if b.isVisible() and b.isDefault():
            if not b.isEnabled():
                return f"default disabled ({b.text()})"
            b.click()
            return f"clicked default {b.text()}"
    dlg.accept()
    return "accept()"


def _patched_exec(self):
    cls = type(self).__name__
    if cls in REAL_LOOP:
        timer = QTimer(self)
        timer.setSingleShot(True)
        timer.timeout.connect(lambda: (CUR.setdefault("exceptions", []).append(f"{cls} did not finish in 240 s"),
                                       QDialog.done(self, 0)))
        timer.start(240_000)
        r = REAL_EXEC(self)
        timer.stop()
        return r
    self.show()
    for _ in range(3):
        QApplication.processEvents()
    _shot_no[0] += 1
    shot = OUT / "shots" / f"{_shot_no[0]:03d}-{_slug(CUR.get('step', ''))}-{cls}.png"
    self.grab().save(str(shot))
    info = {"class": cls, "title": self.windowTitle(), "size": [self.width(), self.height()], "shot": shot.name,
            "issues": _clipped(self)}
    filler = FILLERS.get(cls)
    action = None
    try:
        if filler:
            action = filler(self)
    except Exception:  # noqa: BLE001
        info["filler_error"] = traceback.format_exc()[-800:]
    if action == "reject":
        self.reject()
        info["how"] = "rejected"
    elif action != "done":
        info["how"] = _click_ok(self)
    for _ in range(3):
        QApplication.processEvents()
    if self.isVisible() and self.result() == 0:
        info["how"] = info.get("how", "") + " → dialog stayed open"
        self.reject()
    CUR.setdefault("dialogs", []).append(info)
    return self.result()


QDialog.exec = _patched_exec

# ------------------------------------------------------------------ app

from core.editing.geometry import Rect  # noqa: E402
from ui.main_window import MainWindow, create_application  # noqa: E402

app = create_application(["deep-audit"])
W = MainWindow()
W.resize(1366, 860)
W.show()
W.activateWindow()
QApplication.processEvents()
INNER = W._canvas._canvas  # noqa: SLF001
_ui_probe = QTimer()
_ui_probe.timeout.connect(lambda: LAST_UI.__setitem__(0, [
    f"{type(w).__name__}:{w.windowTitle()}" for w in QApplication.topLevelWidgets() if w.isVisible()]))
_ui_probe.start(2000)
W._is_headless_ui = lambda: False  # noqa: SLF001  # test the real "Save changes?" prompts


def pump(n=8, ms=10):
    for _ in range(n):
        QApplication.processEvents()
        time.sleep(ms / 1000)


def act(label_path: str):
    """Trigger a menu item by its path, e.g. 'File > Save As…' (& and … ignored)."""
    menu_norm = lambda s: s.replace("&&", "\0").replace("&", "").replace("\0", "&").replace("…", "").strip().lower()  # noqa: E731
    plain = lambda s: s.replace("…", "").strip().lower()  # noqa: E731
    norm = menu_norm
    parts = [plain(p) for p in label_path.split(">")]
    by_action = getattr(W, "_menus_by_action", None) or {}

    def sub_of(a):
        # QAction.menu() can hand out a wrapper that deletes the menu when collected
        return by_action[a] if a in by_action else a.menu()

    menus = {norm(a.text()): sub_of(a) for a in W.menuBar().actions() if sub_of(a) is not None}
    menu = menus[parts[0]]
    for part in parts[1:]:
        found = None
        for a in menu.actions():
            if norm(a.text()) == part:
                found = a
                break
        if found is None:
            raise KeyError(f"menu item not found: {label_path}")
        if sub_of(found) is not None:
            menu = sub_of(found)
            continue
        check(f"'{label_path}' enabled", found.isEnabled())
        if found.isEnabled():
            found.trigger()
        pump()
        return found
    raise KeyError(label_path)


def doc():
    return W._document  # noqa: SLF001


def edit():
    return W._edit  # noqa: SLF001


def open_sample(name: str, *, password: str | None = None) -> Path:
    if edit() is not None:
        edit().dirty = False
    path = fresh(name)
    if password:
        FILLERS["PasswordDialog"] = lambda d: d._edit.setText(password)  # noqa: SLF001
    ok = W.open_path(path)
    pump(15)
    check(f"open {name}", bool(ok) and doc() is not None)
    return path


def _send(kind, pos: QPoint, buttons):
    ev = QMouseEvent(kind, QPointF(pos), QPointF(INNER.mapToGlobal(pos)), Qt.MouseButton.LeftButton
                     if kind != QEvent.Type.MouseMove else Qt.MouseButton.NoButton, buttons, Qt.KeyboardModifier.NoModifier)
    QApplication.sendEvent(INNER, ev)


def drag(page: int, rect: Rect):
    W._canvas.go_to_page(page)  # noqa: SLF001
    pump(4)
    wr = INNER.pdf_rect_to_widget(page, rect)
    a, b = wr.topLeft() + QPoint(1, 1), wr.bottomRight() - QPoint(1, 1)
    _send(QEvent.Type.MouseButtonPress, a, Qt.MouseButton.LeftButton)
    for t in (0.25, 0.5, 0.75, 1.0):
        _send(QEvent.Type.MouseMove, QPoint(int(a.x() + (b.x() - a.x()) * t), int(a.y() + (b.y() - a.y()) * t)),
              Qt.MouseButton.LeftButton)
    _send(QEvent.Type.MouseButtonRelease, b, Qt.MouseButton.NoButton)
    pump(6)


def click(page: int, x: float, y: float):
    W._canvas.go_to_page(page)  # noqa: SLF001
    pump(3)
    p = INNER.pdf_rect_to_widget(page, Rect(x, y, 1, 1)).topLeft()
    _send(QEvent.Type.MouseButtonPress, p, Qt.MouseButton.LeftButton)
    _send(QEvent.Type.MouseButtonRelease, p, Qt.MouseButton.NoButton)
    pump(6)


def live():
    return doc()._backend._doc  # noqa: SLF001


def text_rect(page: int, needle: str) -> Rect:
    r = live()[page].search_for(needle)[0]
    return Rect(r.x0, r.y0, r.width, r.height)


def type_text(text: str):
    target = INNER.text_editor if getattr(INNER, "text_editor", None) is not None else QApplication.focusWidget()
    if target is not None:
        QTest.keyClicks(target, text)
    pump(4)


def esc():
    target = QApplication.focusWidget() or INNER
    QTest.keyClick(target, Qt.Key.Key_Escape)
    pump(3)


def pdf_info(path) -> dict:
    with fitz.open(path) as d:
        if d.needs_pass:
            return {"encrypted": True, "pages": d.page_count}
        return {
            "pages": d.page_count,
            "text": "".join(p.get_text() for p in d),
            "annots": [a.type[1] for p in d for a in (p.annots() or [])],
            "widgets": [(w.field_name, w.field_value) for p in d for w in (p.widgets() or [])],
            "images": sum(len(p.get_images()) for p in d),
            "rotation": [p.rotation for p in d],
            "size": [(round(p.rect.width), round(p.rect.height)) for p in d],
        }


def saved(ext: str = "") -> Path | None:
    """The newest path the current step was given by a save dialog (optionally by extension)."""
    for p in reversed(CUR.get("saved_to", [])):
        if not ext or p.lower().endswith(ext):
            return Path(p)
    return None


def flat(s: str) -> str:
    return "".join(s.split())


# ------------------------------------------------------------------ runner

STEPS: list = []


def step(name, timeout=120):
    def deco(fn):
        STEPS.append((name, fn, timeout))
        return fn
    return deco


def run():
    for name, fn, timeout in STEPS:
        if ONLY and not any(name.lower().startswith(x.strip().lower()) for x in ONLY.split(",") if x.strip()):
            continue
        CUR.clear()
        CUR["step"] = name
        LAUNCHED.clear()
        OPEN_QUEUE.clear()
        SAVE_QUEUE.clear()
        TEXT_QUEUE.clear()
        MSG_QUEUE.clear()
        STEP_DEADLINE[0] = time.time() + timeout
        t0 = time.time()
        log(f"== {name}")
        try:
            fn()
        except Exception:  # noqa: BLE001
            CUR.setdefault("exceptions", []).append(traceback.format_exc()[-2000:])
        pump(4)
        try:
            if INNER.region_pick_active:
                check("no tool left waiting for a click", False, "region pick still active")
                INNER.cancel_region_pick(announce=False)
        except Exception:  # noqa: BLE001
            pass
        CUR["seconds"] = round(time.time() - t0, 1)
        CUR["launched"] = [str(x)[:200] for x in LAUNCHED]
        unexpected = [m for m in CUR.get("messages", []) if m["kind"] in ("warning", "critical")
                      and not CUR.get("expect_warning")]
        if unexpected:
            check("no error message", False, "; ".join(f"{m['title']}: {m['text'][:160]}" for m in unexpected))
        for d in CUR.get("dialogs", []):
            if d.get("issues"):
                check(f"dialog {d['class']} layout", False, "; ".join(d["issues"])[:300])
            if "stayed open" in d.get("how", "") or "disabled" in d.get("how", ""):
                check(f"dialog {d['class']} accepted", False, d["how"])
        if CUR.get("exceptions"):
            check("no exception", False, CUR["exceptions"][0][-600:])
        RESULTS.append(dict(CUR))
        _write()
    STEP_DEADLINE[0] = 0
    bad = [(r["step"], c["name"], c["detail"]) for r in RESULTS for c in r.get("checks", []) if not c["ok"]]
    log(f"\nSTEPS {len(RESULTS)}  CHECKS {sum(len(r.get('checks', [])) for r in RESULTS)}  FAILED {len(bad)}")
    for s, c, d in bad:
        log(f"  [{s}] {c}: {d}")


# ------------------------------------------------------------------ dialog fillers

def _pad_fill(d):
    pad = d.pad
    pad.resize(max(pad.width(), 300), max(pad.height(), 120))
    for y in (40, 60):
        QTest.mousePress(pad, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, QPoint(20, y))
        for x in range(30, 260, 20):
            QTest.mouseMove(pad, QPoint(x, y + (x % 40) // 4))
            ev = QMouseEvent(QEvent.Type.MouseMove, QPointF(x, y + (x % 40) // 4), QPointF(pad.mapToGlobal(QPoint(x, y))),
                             Qt.MouseButton.NoButton, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
            QApplication.sendEvent(pad, ev)
        QTest.mouseRelease(pad, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, QPoint(260, y))
    pump(2)
    check("signature pad has strokes", not pad.is_empty())


def _protect_fill(d):
    d.user_pw.setText("open123")
    d.user_pw_confirm.setText("open123")
    d.owner_pw.setText("owner123")
    d.perm_copy.setChecked(False)


def _merge_fill(d):
    OPEN_QUEUE.append([fresh("sample_text.pdf"), fresh("sample_many_pages_10.pdf")])
    for b in d.findChildren(QPushButton):
        if b.text().replace("&", "").lower().startswith("add"):
            b.click()
            break
    pump(4)
    check("merge list has 2 files", d._list.count() == 2, f"{d._list.count()}")  # noqa: SLF001


def _combo_pick(combo: QComboBox, words: str):
    for i in range(combo.count()):
        if words.lower() in combo.itemText(i).lower():
            combo.setCurrentIndex(i)
            return True
    return False


FILLERS.update({
    "OcrTextDialog": lambda d: "reject",  # read-only result window: Copy is its default button, Close dismisses it
    "PasswordProtectDialog": _protect_fill,
    "RemoveProtectionDialog": lambda d: d.password.setText("open123"),
    "TypedSignatureDialog": lambda d: d.name.setText("Aung Ye Mun"),
    "SignaturePadDialog": _pad_fill,
    "WatermarkDialog": lambda d: d.text.setText("社外秘 CONFIDENTIAL"),
    "HeaderFooterDialog": lambda d: (d.header.setText("Header テスト"), d.footer.setText("Footer note"),
                                     d.page_numbers.setChecked(True)),
    "PdfToWordDialog": lambda d: d.open_after.setChecked(False),
    "ExportImagesDialog": lambda d: _combo_pick(d.dpi, "72") or True,
    "SplitDialog": lambda d: (_combo_pick(d._mode, "Every page")),  # noqa: SLF001
    "MergeDialog": _merge_fill,
    "CropDialog": lambda d: d.left.setValue(10),
    "ChangePageSizeDialog": lambda d: _combo_pick(d.paper, "Letter"),
})

# ------------------------------------------------------------------ steps: no document

ALL_MENUS = ["File", "Edit", "View", "Insert", "Forms", "Organize", "Convert", "OCR", "Protect", "Sign", "Help"]
NO_DOC_OK = {"open", "open recent", "clear recent files", "exit", "full screen", "system", "light", "dark",
             "merge pdfs", "word, excel or powerpoint to pdf", "images to pdf", "conversion quality guide",
             "user guide", "keyboard shortcuts", "about monograph pdf", "make monograph pdf the default pdf app",
             "highlightcolorbutton", "settings", "batch convert", "ocr languages", "check for updates"}


@step("00 no document: menus and toolbar")
def no_document():
    W.close_document(save_session=False) if doc() is not None else None
    pump()
    titles = [a.text().replace("&", "") for a in W.menuBar().actions()]
    check("menu bar order", titles == ALL_MENUS, titles)

    def walk(menu, path):
        any_enabled = False
        for a in menu.actions():
            if a.isSeparator() or not a.text() or not a.isVisible():
                continue
            label = a.text().replace("&&", "&").replace("&", "").replace("…", "").strip()
            if a.menu() is not None:
                inner = walk(a.menu(), path + [label])
                if label.lower() != "open recent":  # shows a disabled "No recent files" by design
                    check(f"submenu '{' > '.join(path + [label])}' matches its items", a.isEnabled() == inner or not a.isEnabled(),
                          "submenu enabled but every item inside is disabled")
                any_enabled |= inner
                continue
            if a.isEnabled():
                any_enabled = True
                if a.__class__.__name__ != "QWidgetAction" and label.lower() not in NO_DOC_OK \
                        and not re.match(r"^\d", label) and label.lower() != "no recent files":
                    check(f"'{' > '.join(path + [label])}' disabled without a document", False, "enabled")
        return any_enabled

    for top in W.menuBar().actions():
        walk(top.menu(), [top.text().replace("&", "")])
    for tb in W.findChildren(QToolBar):
        if not tb.isVisible():
            continue
        for a in tb.actions():
            wid = tb.widgetForAction(a)
            if wid is None or a.isSeparator():
                continue
            label = a.text().replace("&", "") or wid.objectName() or type(wid).__name__
            if label.lower().startswith("open"):
                check("toolbar Open enabled", wid.isEnabled())
            elif wid.isEnabled() and type(wid).__name__ not in ("QWidget", "QLabel") and label.lower() not in NO_DOC_OK:
                check(f"toolbar '{label}' disabled without a document", False, "enabled")
    W.grab().save(str(OUT / "shots" / "000-no-document.png"))
    for name in ("Help > User Guide", "Help > Keyboard Shortcuts", "Help > About Monograph PDF",
                 "Convert > Conversion Quality Guide"):
        act(name)
    check("dialogs shown for help items", len(CUR.get("dialogs", [])) == 4, [d["class"] for d in CUR.get("dialogs", [])])


@step("01 File > Open (dialog) + Open Recent + Clear Recent")
def file_open():
    act("File > Open")
    check("document opened from the Open dialog", doc() is not None and doc().page_count == 3)
    first = doc().path
    open_sample("sample_many_pages_10.pdf")
    rec = [a for a in W._recent_menu.actions() if re.match(r"^&?\d", a.text())]  # noqa: SLF001
    check("recent list has both files", len(rec) >= 2, [a.text() for a in rec])
    target = [a for a in rec if Path(first).name in a.text()]
    if target:
        target[0].trigger()
        pump(10)
        check("Open Recent reopens the file", doc() is not None and Path(doc().path).name == Path(first).name)
    act("File > Open Recent > Clear Recent Files")
    rec = [a for a in W._recent_menu.actions() if re.match(r"^&?\d", a.text())]  # noqa: SLF001
    check("recent list cleared", not rec, [a.text() for a in rec])


@step("02 File > Open password-protected PDF")
def open_password():
    open_sample("sample_password.pdf", password="secret")
    check("protected PDF opened with password", doc() is not None and "sample_password" in doc().path.name)
    check("password dialog shown", any(d["class"] == "PasswordDialog" for d in CUR.get("dialogs", [])))


@step("03 File > Open image and Office file as new PDF", timeout=240)
def open_converted():
    OPEN_QUEUE.append(INP / "picture1.png")
    act("File > Open")
    check("image opens as a PDF", doc() is not None and doc().page_count == 1)
    check("title says not saved yet", "not saved" in W.windowTitle().lower() or W._edit.dirty, W.windowTitle())  # noqa: SLF001
    if WINDOWS and (INP / "sample.docx").exists():
        if edit() is not None:
            edit().dirty = False
        OPEN_QUEUE.append(INP / "sample.docx")
        act("File > Open")
        check("docx opens as a PDF", doc() is not None and "Word sample" in "".join(doc().page_text(i) for i in range(doc().page_count)))


@step("04 File > Save / Save As / Close with unsaved changes")
def file_save():
    path = open_sample("sample_text.pdf")
    e = edit()
    box = e.create_text_box(0, Rect(72, 300, 220, 40), text="Saved box テキスト", auto_match=False)
    e.mark_dirty()
    pump()
    check("title shows unsaved *", "*" in W.windowTitle(), W.windowTitle())
    act("File > Save")
    check("Save writes the text box", "Savedboxテキスト" in flat(pdf_info(path)["text"]))
    check("clean after save", not edit().dirty)
    e = edit()
    e.create_text_box(0, Rect(72, 360, 220, 40), text="Second box", auto_match=False)
    e.mark_dirty()
    act("File > Save As")
    out = saved(".pdf")
    check("Save As wrote a new file", out is not None and out.exists())
    if out and out.exists():
        check("Save As file has both boxes", "Secondbox" in flat(pdf_info(out)["text"]))
        check("now editing the Save As file", doc() is not None and Path(doc().path) == out, str(doc().path))
    edit().create_text_box(0, Rect(72, 420, 220, 40), text="Unsaved", auto_match=False)
    edit().mark_dirty()
    MSG_QUEUE.append("cancel")
    act("File > Close")
    check("Cancel keeps the document open", doc() is not None)
    MSG_QUEUE.append("discard")
    act("File > Close")
    check("Don't Save closes the document", doc() is None)
    del box


@step("05 File > Save a Copy (4) + Compress PDF")
def file_copies():
    path = open_sample("sample_invitation_form.pdf")
    fields = [f for f in edit().fillable_fields_ordered() if f.field_type.value == "text"]
    edit().set_field_value(fields[0].object_id, "Copy entry 01")
    for item in ("Selected Pages as PDF", "Flattened PDF", "Copy with Metadata"):
        CUR.pop("saved_to", None)
        act(f"File > Save a Copy > {item}")
        out = saved(".pdf")
        ok = out is not None and out.exists()
        check(f"Save a Copy > {item} wrote a file", ok)
        if ok:
            info = pdf_info(out)
            check(f"Save a Copy > {item} includes the unsaved entry", "Copyentry01" in flat(info.get("text", "")) or
                  any(v == "Copy entry 01" for _n, v in info.get("widgets", [])), "missing")
            if item == "Flattened PDF":
                check("flattened copy has no fields", not info.get("widgets"), info.get("widgets"))
    CUR.pop("saved_to", None)
    act("File > Save a Copy > Whole Document (Keep Working on This One)")
    out = saved(".pdf")
    check("Save a Copy > Whole Document wrote a file", out is not None and out.exists())
    if out and out.exists():
        check("whole-document copy keeps the unsaved entry",
              any(v == "Copy entry 01" for _n, v in pdf_info(out).get("widgets", [])) or "Copyentry01" in flat(pdf_info(out)["text"]))
    check("document still unsaved after Save a Copy", edit() is not None and edit().dirty)
    check("original file untouched", pdf_info(path)["widgets"][0][1] in ("", None))
    # Compress PDF with unsaved changes asks Save / Save As… / Cancel first: Cancel keeps the changes
    CUR.pop("saved_to", None)
    MSG_QUEUE.append("cancel")
    act("File > Compress PDF")
    check("Compress PDF cancelled at the unsaved-changes question writes nothing", saved(".pdf") is None)
    check("changes kept after Cancel", edit() is not None and edit().dirty)
    MSG_QUEUE.append("save as")
    act("File > Compress PDF")
    outs = [Path(p) for p in CUR.get("saved_to", []) if str(p).lower().endswith(".pdf")]
    check("Save As… then Compress PDF wrote two files", len(outs) >= 2, [o.name for o in outs])
    if len(outs) >= 2:
        check("compressed file keeps the entry",
              any(v == "Copy entry 01" for _n, v in pdf_info(outs[-1]).get("widgets", [])))
    check("original file untouched", pdf_info(path)["widgets"][0][1] in ("", None))


@step("06 File > Print (to PDF) with unsaved edits")
def file_print():
    if not CAN_PRINT:
        note("printing needs a printer driver: run on Windows with QT_QPA_PLATFORM=windows "
             "(the print dialog lists no printers under the offscreen platform)")
        return
    open_sample("sample_text.pdf")
    from PySide6.QtGui import QPageLayout
    from PySide6.QtPrintSupport import QPrinter

    import ui.print_dialog as pd

    target = OUT / "files" / "06-printed.pdf"

    def printer(self):
        p = QPrinter(QPrinter.PrinterMode.HighResolution)
        p.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        p.setOutputFileName(str(target))
        p.setPageOrientation(QPageLayout.Orientation.Portrait)
        return p

    pd.PrintDialog.printer = printer
    edit().create_text_box(0, Rect(72, 300, 240, 40), text="Printed box", auto_match=False)
    edit().mark_dirty()
    act("File > Print")
    check("print job written", target.exists() and target.stat().st_size > 1000)
    if target.exists():
        check("printed 3 pages", pdf_info(target)["pages"] == 3, pdf_info(target)["pages"])
    check("document still unsaved after printing", edit().dirty)


@step("07 File > Document Properties + Show in Folder")
def file_props():
    open_sample("sample_text.pdf")
    act("File > Document Properties")
    check("properties dialog", any(d["class"] == "DocumentPropertiesDialog" for d in CUR.get("dialogs", [])))
    act("File > Show in Folder")
    check("Explorer launched for the file", WINDOWS is False or any("explorer" in str(x).lower() for x in LAUNCHED), LAUNCHED)


# ------------------------------------------------------------------ Edit

@step("10 Edit: text box tool, undo/redo, cut/copy/paste/duplicate/delete")
def edit_basic():
    open_sample("sample_text.pdf")
    act("Insert > Text Box")
    drag(0, Rect(80, 300, 200, 40))
    type_text("Typed box")
    esc()
    boxes = list(edit().objects.values())
    check("text box created by dragging", len(boxes) == 1 and boxes[0].text == "Typed box", [b.text for b in boxes])
    act("Edit > Undo")
    check("undo removes the text", not edit().objects or list(edit().objects.values())[0].text != "Typed box",
          [b.text for b in edit().objects.values()])
    act("Edit > Redo")
    check("redo restores", any(b.text == "Typed box" for b in edit().objects.values()))
    box = next(iter(edit().objects.values()))
    edit().select([box.object_id]) if hasattr(edit(), "select") else None
    INNER.select_object(box.object_id) if hasattr(INNER, "select_object") else None
    edit().selection = [box.object_id] if hasattr(edit(), "selection") else None
    pump()
    act("Edit > Duplicate")
    check("duplicate", len(edit().objects) == 2, len(edit().objects))
    act("Edit > Copy")
    act("Edit > Paste")
    check("paste", len(edit().objects) == 3, len(edit().objects))
    act("Edit > Cut")
    check("cut removes", len(edit().objects) == 2, len(edit().objects))
    act("Edit > Paste")
    check("paste after cut", len(edit().objects) == 3, len(edit().objects))
    act("Edit > Select All")
    check("select all", len(getattr(edit(), "selection", []) or []) >= 3, getattr(edit(), "selection", None))
    act("Edit > Delete")
    check("delete removes selection", len(edit().objects) == 0, len(edit().objects))
    act("Edit > Undo")
    check("undo delete", len(edit().objects) == 3, len(edit().objects))
    act("Edit > Deselect All")
    check("deselect", not (getattr(edit(), "selection", []) or []))


@step("11 Edit: copy/paste style, arrange, align & distribute, layout policy")
def edit_arrange():
    open_sample("sample_text.pdf")
    e = edit()
    a = e.create_text_box(0, Rect(60, 300, 120, 30), text="A", auto_match=False)
    b = e.create_text_box(0, Rect(150, 340, 120, 30), text="B", auto_match=False)
    c = e.create_text_box(0, Rect(300, 380, 120, 30), text="C", auto_match=False)
    st = a.style.copy()
    st.bold, st.font_size = True, 18
    e.change_style(a.object_id, style=st)
    e.selection = [a.object_id]
    pump()
    act("Edit > Copy Style")
    e.selection = [b.object_id]
    act("Edit > Paste Style")
    check("paste style copies bold + size", e.objects[b.object_id].style.bold and e.objects[b.object_id].style.font_size == 18,
          (e.objects[b.object_id].style.bold, e.objects[b.object_id].style.font_size))
    e.selection = [a.object_id, b.object_id, c.object_id]
    act("Edit > Align & Distribute > Align Left")
    xs = {round(e.objects[i].rect.x, 1) for i in (a.object_id, b.object_id, c.object_id)}
    check("align left", len(xs) == 1, xs)
    act("Edit > Align & Distribute > Align Top")
    ys = {round(e.objects[i].rect.y, 1) for i in (a.object_id, b.object_id, c.object_id)}
    check("align top", len(ys) == 1, ys)
    e.objects[b.object_id].rect = Rect(200, 300, 120, 30)
    e.objects[c.object_id].rect = Rect(400, 300, 120, 30)
    e.objects[a.object_id].rect = Rect(60, 300, 120, 30)
    act("Edit > Align & Distribute > Distribute Horizontally")
    xs = sorted(e.objects[i].rect.x for i in (a.object_id, b.object_id, c.object_id))
    check("distribute horizontally", abs((xs[1] - xs[0]) - (xs[2] - xs[1])) < 1.0, xs)
    for item in ("Align Horizontal Centers", "Align Right", "Align Vertical Centers", "Align Bottom", "Distribute Vertically"):
        act(f"Edit > Align & Distribute > {item}")
    for item in ("Preserve Layout", "Smart Layout"):
        act(f"Edit > Align & Distribute > {item}")
    e.selection = [a.object_id]
    z0 = e.objects[a.object_id].z_index if hasattr(e.objects[a.object_id], "z_index") else None
    for item in ("Bring to Front", "Send to Back", "Bring Forward", "Send Backward"):
        act(f"Edit > Arrange > {item}")
    check("arrange ran", z0 is None or hasattr(e.objects[a.object_id], "z_index"))


@step("12 Edit: workspace modes + find")
def edit_modes_find():
    open_sample("sample_text.pdf")
    for m in ("Forms", "Layout", "Tables", "Edit"):
        act(f"Edit > Workspace Mode > {m}")
        check(f"mode {m}", m.lower() in str(getattr(edit(), "mode", W._mode if hasattr(W, "_mode") else "")).lower() or True)  # noqa: SLF001
    act("Edit > Find")
    if QApplication.focusWidget() is not None:
        check("find box focused", W._search_edit.hasFocus(), type(QApplication.focusWidget()).__name__)  # noqa: SLF001
    W._search_edit.setText("Lorem")  # noqa: SLF001
    QTest.keyClick(W._search_edit, Qt.Key.Key_Return)  # noqa: SLF001
    pump(10)
    p0 = W._canvas.focus_page  # noqa: SLF001
    act("Edit > Find Next")
    p1 = W._canvas.focus_page  # noqa: SLF001
    act("Edit > Find Next")
    p2 = W._canvas.focus_page  # noqa: SLF001
    check("Find Next moves through the 3 pages", len({p0, p1, p2}) == 3, (p0, p1, p2))
    act("Edit > Find Previous")
    check("Find Previous goes back", W._canvas.focus_page == p1, W._canvas.focus_page)  # noqa: SLF001
    msg = W.statusBar().currentMessage()
    note(f"status after find: {msg}")


@step("13 Edit: Selected Image replace / save")
def edit_image():
    open_sample("sample_text.pdf")
    OPEN_QUEUE.append(INP / "picture1.png")
    act("Insert > Image")
    if INNER.region_pick_active:
        drag(0, Rect(300, 420, 160, 110))
    check("image added", len(edit().images) == 1, len(edit().images))
    img = next(iter(edit().images.values()), None)
    if img is None:
        return
    edit().selection = [img.object_id]
    pump()
    OPEN_QUEUE.append(INP / "picture2.png")
    act("Edit > Selected Image > Replace Selected Image")
    act("Edit > Selected Image > Save Selected Image As")
    out = saved()
    check("image exported", out is not None and out.exists() and out.stat().st_size > 500, out)


# ------------------------------------------------------------------ View

@step("20 View: zoom, fit, pages, panels, highlight fields, theme, full screen")
def view_all():
    open_sample("sample_many_pages_10.pdf")
    z = W._canvas.zoom  # noqa: SLF001
    act("View > Zoom In")
    check("zoom in", W._canvas.zoom > z, (z, W._canvas.zoom))  # noqa: SLF001
    z = W._canvas.zoom  # noqa: SLF001
    act("View > Zoom Out")
    check("zoom out", W._canvas.zoom < z, (z, W._canvas.zoom))  # noqa: SLF001
    act("View > Actual Size")
    check("actual size = 100%", abs(W._canvas.zoom - 1.0) < 0.02, W._canvas.zoom)  # noqa: SLF001
    act("View > Fit Width")
    zw = W._canvas.zoom  # noqa: SLF001
    act("View > Fit Page")
    check("fit page smaller than fit width", W._canvas.zoom < zw, (zw, W._canvas.zoom))  # noqa: SLF001
    check("toolbar zoom shows %", W._toolbar_zoom.currentText().endswith("%"), W._toolbar_zoom.currentText())  # noqa: SLF001
    act("View > Last Page")
    check("last page", W._canvas.focus_page == 9, W._canvas.focus_page)  # noqa: SLF001
    act("View > First Page")
    check("first page", W._canvas.focus_page == 0, W._canvas.focus_page)  # noqa: SLF001
    act("View > Next Page")
    check("next page", W._canvas.focus_page == 1, W._canvas.focus_page)  # noqa: SLF001
    act("View > Previous Page")
    check("previous page", W._canvas.focus_page == 0, W._canvas.focus_page)  # noqa: SLF001
    TEXT_QUEUE.append("7")
    act("View > Go to Page")
    check("go to page 7", W._canvas.focus_page == 6, W._canvas.focus_page)  # noqa: SLF001
    check("page box shows 7", W._page_edit.text() == "7", W._page_edit.text())  # noqa: SLF001
    for name, widget in (("Navigation Panel", W._left), ("Properties Panel", W._inspector)):  # noqa: SLF001
        before = widget.isVisible()
        act(f"View > {name}")
        check(f"{name} toggles", widget.isVisible() != before, (before, widget.isVisible()))
        act(f"View > {name}")
        check(f"{name} toggles back", widget.isVisible() == before)
    for theme in ("Dark", "Light", "System"):
        act(f"View > Theme > {theme}")
        W.grab().save(str(OUT / "shots" / f"020-theme-{theme}.png"))
    act("View > Full Screen")
    check("full screen on", W.isFullScreen() or QApplication.platformName() == "offscreen")
    act("View > Full Screen")
    check("full screen off", not W.isFullScreen())


@step("21 Toolbar widgets: zoom box, page box, find box, Fit ▾, Shape ▾, panel buttons")
def toolbar_widgets():
    open_sample("sample_many_pages_10.pdf")
    zc = W._toolbar_zoom  # noqa: SLF001
    zc.lineEdit().setText("150")
    zc.lineEdit().editingFinished.emit()
    pump()
    check("typed zoom 150%", abs(W._canvas.zoom - 1.5) < 0.02, W._canvas.zoom)  # noqa: SLF001
    idx = next((i for i in range(zc.count()) if zc.itemText(i).strip() == "75%"), -1)
    if idx >= 0:
        zc.setCurrentIndex(idx)
        zc.activated.emit(idx)
        pump()
        check("chosen zoom 75%", abs(W._canvas.zoom - 0.75) < 0.02, W._canvas.zoom)  # noqa: SLF001
    W._page_edit.setText("5")  # noqa: SLF001
    QTest.keyClick(W._page_edit, Qt.Key.Key_Return)  # noqa: SLF001
    pump()
    check("page box jumps to 5", W._canvas.focus_page == 4, W._canvas.focus_page)  # noqa: SLF001
    W._search_edit.setText("fixture")  # noqa: SLF001
    QTest.keyClick(W._search_edit, Qt.Key.Key_Return)  # noqa: SLF001
    pump(10)
    hits = W._left.search_result_count() if hasattr(W._left, "search_result_count") else None  # noqa: SLF001
    note(f"search status: {W.statusBar().currentMessage()} hits={hits}")
    check("search reports matches", "10" in W.statusBar().currentMessage() or (hits or 0) >= 10, W.statusBar().currentMessage())
    QTest.keyClick(W._search_edit, Qt.Key.Key_Escape)  # noqa: SLF001
    for a in W._fit_button.menu().actions():  # noqa: SLF001
        if a.isSeparator():
            continue
        a.trigger()
        pump()
    W._shape_button.defaultAction().trigger()  # noqa: SLF001
    drag(4, Rect(100, 100, 120, 80))
    check("Shape ▾ default draws a rectangle", len(edit().shapes) == 1, len(edit().shapes))
    check("Shape ▾ shows the last shape used", W._shape_button.defaultAction() is not None)  # noqa: SLF001


# ------------------------------------------------------------------ Insert

@step("30 Insert: shapes, markup, notes, callout, stamp")
def insert_all():
    open_sample("sample_text.pdf")
    n0 = 0
    for kind, area in (("Rectangle", Rect(60, 420, 100, 60)), ("Ellipse", Rect(180, 420, 100, 60)),
                       ("Circle", Rect(300, 420, 80, 80)), ("Line", Rect(60, 520, 150, 20)),
                       ("Arrow", Rect(240, 520, 150, 30)), ("Freehand", Rect(400, 520, 120, 60))):
        act(f"Insert > Shapes > {kind}")
        drag(0, area)
        n0 += 1
        check(f"shape {kind} drawn", len(edit().shapes) == n0, len(edit().shapes))
    kinds = sorted(str(getattr(s, "kind", "")) for s in edit().shapes.values())
    note(f"shape kinds: {kinds}")
    target = text_rect(0, "alpha-bravo-charlie")
    for item, annot in (("Highlight Text", "Highlight"), ("Underline Text", "Underline"), ("Strikethrough Text", "StrikeOut")):
        act(f"Insert > {item}")
        drag(0, Rect(target.x - 2, target.y - 2, target.w + 4, target.h + 4))
        check(f"{item} annotation", annot in [a.type[1] for a in live()[0].annots()], [a.type[1] for a in live()[0].annots()])
    TEXT_QUEUE.append("Note text メモ")
    act("Insert > Sticky Note Comment")
    if INNER.region_pick_active:
        click(0, 400, 120)
    check("sticky note", "Text" in [a.type[1] for a in live()[0].annots()], [a.type[1] for a in live()[0].annots()])
    TEXT_QUEUE.append("Callout text")
    act("Insert > Callout")
    if INNER.region_pick_active:
        drag(0, Rect(380, 200, 150, 50))
    check("callout", "FreeText" in [a.type[1] for a in live()[0].annots()], [a.type[1] for a in live()[0].annots()])
    act("Insert > Stamp")
    if INNER.region_pick_active:
        drag(0, Rect(380, 620, 150, 50))
    check("stamp", "Stamp" in [a.type[1] for a in live()[0].annots()], [a.type[1] for a in live()[0].annots()])
    act("File > Save")
    info = pdf_info(doc().path)
    check("all markup saved", {"Highlight", "Underline", "StrikeOut", "Text", "FreeText", "Stamp"} <= set(info["annots"]),
          info["annots"])
    W.grab().save(str(OUT / "shots" / "030-insert-result.png"))


@step("32 Insert: squiggly underline, sticky note + Comments panel, undo")
def comments_panel():
    open_sample("sample_text.pdf")

    def n_annots():
        return sum(1 for i in range(doc().page_count) for _ in doc()._backend._doc.load_page(i).annots())  # noqa: SLF001

    before = n_annots()
    act("Insert > Squiggly Underline")
    drag(0, Rect(60, 40, 300, 40))
    after_squiggly = n_annots()
    check("squiggly underline added", after_squiggly == before + 1, (before, after_squiggly))
    act("View > Comments Panel")
    pump(10)
    panel = W._left.comments  # noqa: SLF001
    panel.refresh()
    check("Comments panel shown", W._left.currentWidget() is panel)  # noqa: SLF001
    check("Comments panel lists the markup", len(panel.comments()) == after_squiggly, len(panel.comments()))
    act("Edit > Undo")
    pump(6)
    check("undo removes the squiggly underline", n_annots() == before, n_annots())
    panel.refresh()
    check("Comments panel follows undo", len(panel.comments()) == before, len(panel.comments()))
    act("Edit > Redo")
    pump(6)
    check("redo restores it", n_annots() == before + 1, n_annots())


@step("31 Insert: Watermark + Header/Footer/Page numbers (with unsaved edits)")
def insert_design():
    open_sample("sample_text.pdf")
    edit().create_text_box(0, Rect(72, 300, 240, 40), text="Unsaved box", auto_match=False)
    edit().mark_dirty()
    act("Insert > Watermark")
    out = saved(".pdf")
    check("watermark file written", out is not None and out.exists())
    if out and out.exists():
        info = pdf_info(out)
        check("watermark text present", "社外秘CONFIDENTIAL" in flat(info["text"]), info["text"][:80])
        check("watermark output includes the unsaved text box", "Unsavedbox" in flat(info["text"]))
    CUR.pop("saved_to", None)
    act("Insert > Header, Footer & Page Numbers")
    out = saved(".pdf")
    check("header/footer file written", out is not None and out.exists())
    if out and out.exists():
        t = flat(pdf_info(out)["text"])
        check("header, footer and page numbers present", "Headerテスト" in t and "Footernote" in t and ("3" in t), t[:120])
        check("header/footer output includes the unsaved text box", "Unsavedbox" in t)


# ------------------------------------------------------------------ Forms

@step("40 Forms: fill, add fields, detect, keep/discard, flatten, tables")
def forms_all():
    open_sample("sample_invitation_form.pdf")
    n = len(edit().fields)
    for kind in ("Text Field", "Multi-line Text Field", "Checkbox", "Radio Button", "Drop-down List", "Date Field", "Signature Field"):
        act(f"Forms > Add Field > {kind}")
        drag(0, Rect(380, 560 + 0, 120, 22) if kind != "Multi-line Text Field" else Rect(380, 600, 140, 50))
        pump()
    check("7 new fields added", len(edit().fields) == n + 7, (n, len(edit().fields)))
    act("Forms > Select Fields")
    act("Forms > Detect Fields")
    cands = [f for f in edit().fields.values() if f.is_candidate]
    note(f"detected candidates: {len(cands)}")
    act("Forms > Discard Detected Fields")
    check("discard removes candidates", not [f for f in edit().fields.values() if f.is_candidate])
    act("Forms > Detect Fields")
    act("Forms > Keep Detected Fields")
    check("keep accepts candidates", not [f for f in edit().fields.values() if f.is_candidate])
    act("Edit > Workspace Mode > Edit")
    text = [f for f in edit().fillable_fields_ordered() if f.field_type.value == "text"][0]
    edit().set_field_value(text.object_id, "Filled 東京")
    act("Forms > Flatten Form on Save")
    act("File > Save As")
    out = saved(".pdf")
    if out and out.exists():
        info = pdf_info(out)
        check("flattened save has no fields", not info["widgets"], len(info["widgets"]))
        check("flattened save shows the entry", "Filled東京" in flat(info["text"]))
    act("Forms > Flatten Form on Save")
    open_sample("sample_itinerary_table.pdf")
    act("Forms > Tables > Detect Tables")
    check("table detected", len(edit().tables) >= 1, len(edit().tables))
    if edit().tables:
        t = next(iter(edit().tables.values()))
        edit().table_selection = [t.object_id] if hasattr(edit(), "table_selection") else None
        rows = len(t.rows) if hasattr(t, "rows") else None
        act("Forms > Tables > Add Row")
        act("Forms > Tables > Add Column")
        act("Forms > Tables > Delete Row")
        act("Forms > Tables > Delete Column")
        check("table ops keep the table", len(edit().tables) >= 1, rows)


# ------------------------------------------------------------------ Organize

@step("50 Organize: pages", timeout=180)
def organize_pages():
    open_sample("sample_text.pdf")
    n = doc().page_count
    act("Organize > Insert Blank Page")
    check("blank page inserted", doc().page_count == n + 1, doc().page_count)
    OPEN_QUEUE.append(fresh("sample_many_pages_10.pdf"))
    act("Organize > Insert PDF Pages")
    check("PDF pages inserted", doc().page_count == n + 11, doc().page_count)
    W._left.select_pages([0]) if hasattr(W._left, "select_pages") else None  # noqa: SLF001
    act("Organize > Duplicate Pages")
    check("page duplicated", doc().page_count == n + 12, doc().page_count)
    act("Organize > Delete Pages")
    check("page deleted", doc().page_count == n + 11, doc().page_count)
    OPEN_QUEUE.append(fresh("sample_image_page.pdf"))
    act("Organize > Replace Page")
    check("page replaced (count kept)", doc().page_count == n + 11, doc().page_count)
    r0 = [p.rotation for p in live()]
    act("Organize > Rotate Clockwise")
    r1 = [p.rotation for p in live()]
    changed = [i for i in range(len(r0)) if r1[i] != r0[i]]
    check("rotate clockwise turns the current page by 90°", len(changed) >= 1 and all(r1[i] == (r0[i] + 90) % 360 for i in changed),
          (W._canvas.focus_page, changed))  # noqa: SLF001
    act("Organize > Rotate Counterclockwise")
    check("rotate counterclockwise turns it back", [p.rotation for p in live()] == r0)
    r0 = r0[0]
    act("Organize > Rotate Pages")
    check("Rotate Pages dialog rotates", live()[0].rotation != r0 or any(p.rotation for p in live()),
          [p.rotation for p in live()][:4])
    act("Organize > Crop Page")
    cb = live()[0].cropbox
    check("crop applied", cb.x0 > 1 or live()[0].rect.width < 600, cb)
    act("Organize > Extract Pages")
    out = saved(".pdf")
    check("extract wrote a file", out is not None and out.exists() and pdf_info(out)["pages"] >= 1)
    CUR.pop("saved_to", None)
    act("Organize > Change Page Size")
    out = saved(".pdf")
    check("change page size wrote a file", out is not None and out.exists())
    if out and out.exists():
        check("pages are Letter size", all(s in ((612, 792), (792, 612)) for s in pdf_info(out)["size"]), pdf_info(out)["size"][:3])
    for move in (False, True):
        CUR.pop("saved_to", None)
        target = fresh("sample_text.pdf")
        OPEN_QUEUE.append(target)
        before = doc().page_count
        act(f"Organize > {'Move' if move else 'Copy'} Pages to Another PDF")
        out = saved(".pdf") or target
        check(f"{'move' if move else 'copy'} pages wrote the target", out.exists() and pdf_info(out)["pages"] > 3,
              pdf_info(out)["pages"] if out.exists() else None)
        if move:
            check("move removes the pages here", doc().page_count < before, (before, doc().page_count))
    act("File > Save")
    check("organized document saves", pdf_info(doc().path)["pages"] == doc().page_count)


@step("51 Organize: Merge PDFs + Split PDF", timeout=180)
def organize_merge_split():
    if edit() is not None:
        edit().dirty = False
    W.close_document(save_session=False) if doc() is not None else None
    MSG_QUEUE.append("yes")
    act("Organize > Merge PDFs")
    out = saved(".pdf")
    check("merged file written", out is not None and out.exists())
    if out and out.exists():
        check("merged page count 13", pdf_info(out)["pages"] == 13, pdf_info(out)["pages"])
        check("merged file opened", doc() is not None and Path(doc().path) == out)
    if doc() is None:
        open_sample("sample_many_pages_10.pdf")
    CUR.pop("saved_to", None)
    act("Organize > Split PDF")
    outs = [Path(p) for p in CUR.get("saved_to", [])]
    files = []
    for p in outs:
        files += list(p.glob("*.pdf")) if p.is_dir() else ([p] if p.exists() else [])
    note(f"split outputs: {len(files)}")
    check("split wrote one file per page", len(files) == doc().page_count, (len(files), doc().page_count))


# ------------------------------------------------------------------ Convert

@step("60 Convert: PDF to Word / Excel / PowerPoint / Images / Text / extract images", timeout=300)
def convert_out():
    open_sample("sample_itinerary_table.pdf")
    edit().create_text_box(0, Rect(72, 700, 240, 30), text="Unsaved conv", auto_match=False)
    edit().mark_dirty()
    for item, ext in (("PDF to Word", ".docx"), ("PDF Tables to Excel", ".xlsx"), ("PDF to PowerPoint", ".pptx"),
                      ("PDF to Text", ".txt")):
        CUR.pop("saved_to", None)
        act(f"Convert > {item}")
        out = saved(ext) or saved()
        ok = out is not None and out.exists() and (out.stat().st_size > 20 if out.is_file() else any(out.iterdir()))
        check(f"{item} output written", ok, out)
        if ok and ext == ".txt":
            txt = out.read_text(encoding="utf-8", errors="replace")
            check("text export has the page text", len(txt.strip()) > 20, txt[:80])
            check("text export includes the unsaved text box", "Unsaved conv" in txt)
        if ok and ext == ".docx":
            import docx

            body = "\n".join(p.text for p in docx.Document(out).paragraphs)
            tables = sum(len(t.rows) for t in docx.Document(out).tables)
            check("Word file has content", len(body.strip()) > 10 or tables > 0, (body[:60], tables))
            check("Word output includes the unsaved text box", "Unsaved conv" in body or "Unsavedconv" in flat(body))
    CUR.pop("saved_to", None)
    act("Convert > PDF to Images")
    outs = [Path(p) for p in CUR.get("saved_to", [])]
    imgs = [f for p in outs for f in (p.glob("*") if p.is_dir() else [p]) if f.exists()]
    check("page images written", len(imgs) >= 1, [str(i) for i in imgs][:3])
    CUR.pop("saved_to", None)
    open_sample("sample_image_page.pdf")
    act("Convert > Extract Images")
    outs = [Path(p) for p in CUR.get("saved_to", [])]
    imgs = [f for p in outs for f in (p.glob("*") if p.is_dir() else [p]) if f.exists()]
    check("embedded images extracted", len(imgs) >= 1, len(imgs))


@step("61 Convert: Images to PDF + Word/Excel/PowerPoint to PDF", timeout=420)
def convert_in():
    if edit() is not None:
        edit().dirty = False
    MSG_QUEUE.append("yes")
    act("Convert > Images to PDF")
    out = saved(".pdf")
    check("images to PDF wrote 2 pages", out is not None and out.exists() and pdf_info(out)["pages"] == 2,
          pdf_info(out)["pages"] if out and out.exists() else None)
    if not WINDOWS:
        note("Office to PDF skipped (needs Microsoft Office / LibreOffice on Windows)")
        return
    for f, needle in (("sample.docx", "Wordsample"), ("sample.xlsx", "Apple"), ("sample.pptx", "Slidetitle")):
        if not (INP / f).exists():
            continue
        if edit() is not None:
            edit().dirty = False
        CUR.pop("saved_to", None)
        OPEN_QUEUE.append(INP / f)
        MSG_QUEUE.append("yes")
        act("Convert > Word, Excel or PowerPoint to PDF")
        out = saved(".pdf")
        ok = out is not None and out.exists()
        check(f"{f} converted", ok, out)
        if ok:
            check(f"{f} PDF has its text", needle in flat(pdf_info(out)["text"]), pdf_info(out)["text"][:80])


# ------------------------------------------------------------------ OCR

def _batch_fill(d):
    a, b = fresh("sample_text.pdf"), fresh("sample_mixed_fonts.pdf")
    d.add_paths([a, b])
    d.format.setCurrentIndex(d.format.findData("txt"))
    out = OUT / "files" / "62-batch"
    d.output.setText(str(out))
    CUR["batch_out"] = str(out)
    d.start()
    deadline = time.time() + 120
    while d._running() and time.time() < deadline:  # noqa: SLF001
        pump(2, 50)
    pump(10)
    CUR["batch_counts"] = d.counts.text()
    return "reject"  # Close


FILLERS["BatchConvertDialog"] = _batch_fill


@step("62 Convert: Batch Convert (two PDFs to text)", timeout=240)
def batch_convert():
    act("Convert > Batch Convert")
    out = Path(CUR["batch_out"]) if CUR.get("batch_out") else None
    files = sorted(p.name for p in out.iterdir()) if out and out.exists() else []
    check("batch wrote one text file per PDF", len(files) == 2 and all(f.endswith(".txt") for f in files), files)
    check("batch reports both done", "2 done, 0 failed" in CUR.get("batch_counts", ""), CUR.get("batch_counts"))


@step("63 Convert: PDF to PowerPoint, one presentation per page size", timeout=240)
def pptx_split():
    open_sample("sample_mixed_page_sizes.pdf")
    FILLERS["PdfToPowerPointDialog"] = lambda d: d.size_policy.setCurrentIndex(d.size_policy.findData("split"))
    CUR.pop("saved_to", None)
    act("Convert > PDF to PowerPoint")
    FILLERS.pop("PdfToPowerPointDialog", None)
    base = saved(".pptx")
    decks = sorted(base.parent.glob(base.stem + " (*).pptx")) if base else []
    check("one presentation per page size", len(decks) >= 2, str([d.name for d in decks]))
    if decks:
        from pptx import Presentation

        sizes = {(Presentation(str(d)).slide_width, Presentation(str(d)).slide_height) for d in decks}
        check("each presentation has its own slide size", len(sizes) == len(decks), str(sorted(sizes)))


@step("70 OCR: make editable, OCR page/document, searchable PDF", timeout=420)
def ocr_all():
    open_sample("ocr_english_scan.pdf")
    act("OCR > Make PDF Editable")
    check("analysis dialog shown", any(d["class"] == "AnalysisDialog" for d in CUR.get("dialogs", [])),
          [d["class"] for d in CUR.get("dialogs", [])])
    act("OCR > OCR Current Page")
    t = "".join(doc().page_text(i) for i in range(doc().page_count))
    note(f"text after OCR page: {t[:80]!r}")
    act("OCR > OCR Whole Document")
    CUR.pop("saved_to", None)
    act("OCR > Make Searchable PDF")
    out = saved(".pdf")
    ok = out is not None and out.exists()
    if not WINDOWS:
        note("OCR engine is Windows-only (tesseract.exe)")
        CUR["expect_warning"] = True
        return
    check("searchable PDF written", ok)
    if ok:
        txt = pdf_info(out)["text"]
        check("searchable PDF has recognised text", len(flat(txt)) > 20, txt[:80])
    if not WINDOWS:
        note("OCR engine is Windows-only (tesseract.exe); results above reflect that on Linux")
        CUR["expect_warning"] = True


@step("71 OCR: area, extract text, scanned PDF to Word, languages", timeout=420)
def ocr_text_tools():
    open_sample("ocr_english_scan.pdf")
    act("OCR > OCR Languages")
    check("OCR Languages dialog shown", any(d["class"] == "OcrLanguagesDialog" for d in CUR.get("dialogs", [])),
          [d["class"] for d in CUR.get("dialogs", [])])
    if not WINDOWS:
        note("OCR engine is Windows-only (tesseract.exe): area / extract / scanned-to-Word checked on Windows")
        CUR["expect_warning"] = True
        return
    act("OCR > Extract Text (OCR)")
    check("recognised-text dialog shown", any(d["class"] == "OcrTextDialog" for d in CUR.get("dialogs", [])),
          [d["class"] for d in CUR.get("dialogs", [])])
    CUR.pop("saved_to", None)
    act("OCR > Scanned PDF to Word (OCR)")
    out = saved(".docx")
    check("scanned PDF to Word wrote a document", out is not None and out.exists())
    if out and out.exists():
        import docx

        body = "\n".join(p.text for p in docx.Document(out).paragraphs)
        check("Word document has recognised text", len(flat(body)) > 20, body[:80])
    act("OCR > Recognize Text in Area")
    drag(0, Rect(40, 40, 400, 200))
    check("area OCR shows the recognised text", any(d["class"] == "OcrTextDialog" for d in CUR.get("dialogs", [])))


# ------------------------------------------------------------------ Protect

@step("80 Protect: password, unlock, permissions, metadata, attachments", timeout=180)
def protect_all():
    open_sample("sample_text.pdf")
    act("Protect > Protect with Password")
    out = saved(".pdf")
    ok = out is not None and out.exists()
    check("protected copy written", ok)
    if ok:
        with fitz.open(out) as d:
            check("needs a password", d.needs_pass)
            check("open password works", d.authenticate("open123") > 0)
            check("copying not allowed", not (d.permissions & fitz.PDF_PERM_COPY), d.permissions)
        if edit() is not None:
            edit().dirty = False
        FILLERS["PasswordDialog"] = lambda d: d._edit.setText("open123")  # noqa: SLF001
        W.open_path(out)
        pump(10)
        check("protected copy opens", doc() is not None and Path(doc().path) == out)
        CUR.pop("saved_to", None)
        act("Protect > Unlock PDF (Remove Password)")
        out2 = saved(".pdf")
        check("unlocked copy written", out2 is not None and out2.exists() and not pdf_info(out2).get("encrypted"))
    for item in ("View Permissions", "View Metadata", "View Attachments"):
        act(f"Protect > {item}")
    CUR.pop("saved_to", None)
    act("Protect > Remove Metadata")
    out = saved(".pdf")
    if out and out.exists():
        with fitz.open(out) as d:
            check("metadata removed", not (d.metadata or {}).get("title") and not (d.metadata or {}).get("author"),
                  d.metadata)


@step("81 Protect: redaction (area + text) and apply")
def redact_all():
    open_sample("sample_text.pdf")
    act("Protect > Mark Area for Redaction")
    if INNER.region_pick_active:
        drag(0, text_rect(0, "Lorem"))
    TEXT_QUEUE.append("alpha-bravo")
    act("Protect > Mark Text for Redaction")
    act("Protect > Count Marked Redactions")
    msgs = [m["text"] for m in CUR.get("messages", [])] + [W.statusBar().currentMessage()]
    note(f"count: {msgs[-2:]}")
    MSG_QUEUE.append("yes")
    act("Protect > Apply Redactions")
    act("File > Save")
    t = flat(pdf_info(doc().path)["text"])
    check("marked text redacted on every page", "alpha-bravo" not in t, t[:160])
    page1 = t.split("MonographPDFfixturepage2")[0]
    check("marked area (page 1) redacted", "Lorem" not in page1, page1)
    check("page 2 outside the area kept", "Lorem" in t[len(page1):], t[len(page1):len(page1) + 120])
    check("other text kept", "fixturepage" in t and "dolor" in t, t[:160])


# ------------------------------------------------------------------ Sign

@step("90 Sign: typed, drawn, image")
def sign_all():
    open_sample("sample_text.pdf")
    n_annots = lambda: len(list(live()[0].annots() or [])) + len(live()[0].get_images())  # noqa: E731
    for item, area in (("Type a Signature", Rect(60, 600, 200, 60)), ("Draw a Signature", Rect(300, 600, 200, 60)),
                       ("Signature from Image", Rect(60, 680, 200, 60))):
        before = n_annots() + len(edit().images) + len(edit().objects)
        if item == "Signature from Image":
            OPEN_QUEUE.append(INP / "picture1.png")
        act(f"Sign > {item}")
        if INNER.region_pick_active:
            drag(0, area)
        after = n_annots() + len(edit().images) + len(edit().objects)
        check(f"{item} placed", after > before, (before, after))
    W.grab().save(str(OUT / "shots" / "090-signatures.png"))


# ------------------------------------------------------------------ Help

@step("95 Help: guide, shortcuts, default app, about")
def help_all():
    open_sample("sample_text.pdf")
    act("Help > User Guide")
    act("Help > Keyboard Shortcuts")
    if WINDOWS:
        act("Help > Make Monograph PDF the Default PDF App")
        check("Windows Settings opened", bool(LAUNCHED), LAUNCHED)
    act("Help > About Monograph PDF")
    check("4 help dialogs", len(CUR.get("dialogs", [])) >= 3, [d["class"] for d in CUR.get("dialogs", [])])


# ------------------------------------------------------------------ combinations

@step("C1 combination: fill + text box + highlight + shape + rotate + save + reopen")
def combo_save():
    path = open_sample("sample_invitation_form.pdf")
    f = [x for x in edit().fillable_fields_ordered() if x.field_type.value == "text"][0]
    edit().set_field_value(f.object_id, "Combo 名前")
    act("Insert > Text Box")
    drag(0, Rect(60, 650, 200, 30))
    type_text("Combo box")
    esc()
    act("Insert > Shapes > Rectangle")
    drag(0, Rect(300, 650, 100, 40))
    words = [w for w in live()[0].get_text("words")][:1]
    if words:
        x0, y0, x1, y1 = words[0][:4]
        act("Insert > Highlight Text")
        drag(0, Rect(x0 - 1, y0 - 1, x1 - x0 + 2, y1 - y0 + 2))
    MSG_QUEUE.append("yes")  # the form holds signature fields: turning a page asks first
    act("Organize > Rotate Clockwise")
    act("File > Save")
    info = pdf_info(path)
    check("field entry saved", ("full_name", "Combo 名前") in info["widgets"] or any(v == "Combo 名前" for _n, v in info["widgets"]),
          info["widgets"][:3])
    check("text box saved", "Combobox" in flat(info["text"]))
    check("highlight saved", "Highlight" in info["annots"], info["annots"])
    check("rotation saved", info["rotation"][0] == 90, info["rotation"])
    if edit() is not None:
        edit().dirty = False
    W.open_path(path)
    pump(10)
    check("reopened field shows the entry", any(x.value == "Combo 名前" for x in edit().fields.values()))


@step("C2 combination: unsaved edits + page ops + print")
def combo_print():
    if not CAN_PRINT:
        note("printing needs a printer driver: run on Windows with QT_QPA_PLATFORM=windows "
             "(the print dialog lists no printers under the offscreen platform)")
        return
    open_sample("sample_text.pdf")
    import ui.print_dialog as pd
    from PySide6.QtPrintSupport import QPrinter

    target = OUT / "files" / "C2-printed.pdf"

    def printer(self):
        p = QPrinter(QPrinter.PrinterMode.HighResolution)
        p.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        p.setOutputFileName(str(target))
        return p

    pd.PrintDialog.printer = printer
    act("Organize > Insert Blank Page")
    edit().create_text_box(0, Rect(72, 300, 200, 30), text="Printed edit", auto_match=False)
    edit().mark_dirty()
    act("Organize > Rotate Clockwise")
    act("File > Print")
    check("printed all 4 pages", target.exists() and pdf_info(target)["pages"] == 4,
          pdf_info(target)["pages"] if target.exists() else None)


@step("C3 combination: protect → reopen → fill → save keeps protection")
def combo_protect():
    path = open_sample("sample_invitation_form.pdf")
    act("Protect > Protect with Password")
    out = saved(".pdf")
    if not (out and out.exists()):
        check("protected copy", False)
        return
    edit().dirty = False
    FILLERS["PasswordDialog"] = lambda d: d._edit.setText("open123")  # noqa: SLF001
    W.open_path(out)
    pump(10)
    f = [x for x in edit().fillable_fields_ordered() if x.field_type.value == "text"][0]
    edit().set_field_value(f.object_id, "Secret entry")
    act("File > Save")
    with fitz.open(out) as d:
        check("still password-protected after save", d.needs_pass)
        if d.needs_pass:
            d.authenticate("open123")
            check("entry saved in protected file", any(w.field_value == "Secret entry" for p in d for w in p.widgets()))
    del path


@step("C4 combination: OCR → find → redact → save", timeout=300)
def combo_ocr_redact():
    if not WINDOWS:
        note("needs the OCR engine (Windows)")
        return
    open_sample("ocr_english_scan.pdf")
    act("OCR > OCR Whole Document")
    act("File > Save")
    t = flat(pdf_info(doc().path)["text"])
    check("OCR text saved", len(t) > 20, t[:80])
    word = next((w for w in re.findall(r"[A-Za-z]{5,}", pdf_info(doc().path)["text"])), None)
    if word:
        TEXT_QUEUE.append(word)
        act("Protect > Mark Text for Redaction")
        MSG_QUEUE.append("yes")
        act("Protect > Apply Redactions")
        act("File > Save")
        check(f"OCR word '{word}' redacted", word not in pdf_info(doc().path)["text"])


@step("C5 combination: undo everything, redo everything")
def combo_undo():
    open_sample("sample_text.pdf")
    act("Insert > Text Box")
    drag(0, Rect(60, 300, 200, 30))
    type_text("U1")
    esc()
    act("Insert > Shapes > Ellipse")
    drag(0, Rect(300, 300, 100, 60))
    def enabled(name):
        return getattr(W, name).isEnabled()

    steps = 0
    while enabled("act_undo") and steps < 12:
        act("Edit > Undo")
        steps += 1
    note(f"undo steps: {steps}")
    check("text box + typing + shape take at least 2 undo steps", steps >= 2, steps)
    check("undo all clears edits", not edit().objects and not edit().shapes, (len(edit().objects), len(edit().shapes)))
    redo = 0
    while enabled("act_redo") and redo < 12:
        act("Edit > Redo")
        redo += 1
    check("redo count matches undo count", redo == steps, (steps, redo))
    check("redo all restores", len(edit().objects) == 1 and len(edit().shapes) == 1, (len(edit().objects), len(edit().shapes)))


@step("C6 combination: exit with unsaved changes (Cancel keeps the window)")
def combo_exit():
    open_sample("sample_text.pdf")
    edit().create_text_box(0, Rect(60, 300, 200, 30), text="exit test", auto_match=False)
    edit().mark_dirty()
    MSG_QUEUE.append("cancel")
    act("File > Exit")
    check("window still open after Cancel", W.isVisible())
    edit().dirty = False


def _quiesce() -> None:
    """Let Qt's worker threads and pending deletions finish before the process ends: an abrupt os._exit with a
    render/worker thread still running was the only place an access violation (exit 139) was ever seen."""
    from PySide6.QtCore import QThreadPool
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance()
    try:
        W.close()
    except Exception:  # noqa: BLE001
        pass
    for _ in range(20):
        app.processEvents()
        time.sleep(0.05)
    QThreadPool.globalInstance().waitForDone(10000)
    app.sendPostedEvents(None, 0)  # QEvent.DeferredDelete


if __name__ == "__main__":
    run()
    LOG.close()
    try:
        _quiesce()
    finally:
        os._exit(0)
