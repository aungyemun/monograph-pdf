"""Security regression tests for protected PDFs (0.5.0).

Every operation that wrote an unprotected copy of a password-protected PDF before 0.5.0 must keep
the encryption and permissions; no readable, unencrypted copy of a protected document may be left in
Monograph PDF's temporary folder, also when an operation fails; non-PDF exports are not encrypted
(they carry no PDF security) and that is documented, not hidden."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest

from core.document import open_document

USER, OWNER = "u-secret", "o-secret"
SECRET = "TOPSECRETMARKER"


def _protected(path: Path, *, user: str = USER) -> Path:
    d = fitz.open()
    for i in range(3):
        d.new_page().insert_text((72, 100), f"{SECRET} page {i + 1}", fontsize=16)
    d.set_toc([[1, "Start", 1]])
    d.save(path.as_posix(), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=user, owner_pw=OWNER,
           permissions=fitz.PDF_PERM_PRINT)
    d.close()
    return path


def _check_protected(path: Path, *, open_password: bool = True) -> None:
    with fitz.open(path.as_posix()) as d:
        if open_password:
            assert d.needs_pass, f"{path.name}: open password lost"
            assert d.authenticate(USER) and not d.authenticate("wrong")
        enc = (d.metadata or {}).get("encryption") or ""
        assert "AES" in enc and "256" in enc, f"{path.name}: encryption changed ({enc!r})"
        perms = int(d.permissions)
        assert perms & fitz.PDF_PERM_PRINT and not perms & fitz.PDF_PERM_COPY, f"{path.name}: permissions changed"
        assert not perms & fitz.PDF_PERM_MODIFY
    # the raw bytes never contain the page text in the clear
    assert SECRET.encode() not in path.read_bytes()


def _writers():
    from core.convert.copy_pdf import copy_pdf_preserving_metadata
    from core.convert.flatten_export import export_flattened_pdf
    from core.convert.page_marks import (
        BackgroundSpec,
        HeaderFooterSpec,
        WatermarkSpec,
        apply_page_marks,
        remove_page_marks,
    )
    from core.convert.watermark import apply_page_numbers
    from core.organize.page_ops import extract_pages
    from core.redaction.ops import apply_redactions, mark_region
    from core.security.ops import remove_metadata

    def redact(doc, out):
        from core.editing.geometry import Rect

        mark_region(doc, 0, Rect(60, 80, 100, 30))
        apply_redactions(doc, out)

    def compress(doc, out):
        from core.convert.compress import compress as run, write_result

        write_result(run(doc, mode="lossless", password=USER), out, password=USER)

    return {
        "save_a_copy": copy_pdf_preserving_metadata,
        "remove_metadata": remove_metadata,
        "flattened_pdf": export_flattened_pdf,
        "watermark": lambda d, o: apply_page_marks(d, o, watermark=WatermarkSpec(text="DRAFT")),
        "header_footer": lambda d, o: apply_page_marks(d, o, header_footer=HeaderFooterSpec(header_left="H")),
        "background": lambda d, o: apply_page_marks(d, o, background=BackgroundSpec(color=(0.95, 0.95, 0.9))),
        "page_numbers": apply_page_numbers,
        "remove_page_design": lambda d, o: remove_page_marks(d, o),
        "extract_pages": lambda d, o: extract_pages(d, [0, 2], o),
        "redact": redact,
        "compress": compress,
    }


@pytest.mark.parametrize("name", list(_writers()))
def test_every_previously_affected_output_keeps_encryption_and_permissions(tmp_path, name):
    src = _protected(tmp_path / "in.pdf")
    out = tmp_path / f"{name}.pdf"
    doc = open_document(src, password=USER)
    try:
        _writers()[name](doc, out)
    finally:
        doc.close()
    _check_protected(out)


@pytest.mark.parametrize("name", ["save_a_copy", "watermark", "extract_pages", "remove_metadata"])
def test_permissions_only_protection_is_kept(tmp_path, name):
    src = _protected(tmp_path / "in.pdf", user="")
    out = tmp_path / f"{name}.pdf"
    doc = open_document(src)
    try:
        _writers()[name](doc, out)
    finally:
        doc.close()
    _check_protected(out, open_password=False)


def test_searchable_ocr_output_keeps_protection(tmp_path):
    from core.ocr.engine import OcrUnavailableError, get_ocr_engine
    from core.ocr.pipeline import OcrOptions, run_ocr

    try:
        get_ocr_engine()
    except OcrUnavailableError:
        pytest.skip("OCR engine not available here")
    src = _protected(tmp_path / "in.pdf")
    doc = open_document(src, password=USER)
    try:
        run_ocr(doc, OcrOptions(languages="eng", pages=[0], make_searchable=True, skip_native_text=False),
                destination=tmp_path / "ocr.pdf")
    finally:
        doc.close()
    _check_protected(tmp_path / "ocr.pdf")


# ------------------------------------------------------------------ temporary files


@pytest.fixture
def temp_dir(tmp_path, monkeypatch):
    import app.paths as paths

    t = tmp_path / "runtime-temp"
    monkeypatch.setattr(paths, "TEMP_DIR", t)
    return t


def _leftovers(folder: Path) -> list[Path]:
    return [p for p in folder.rglob("*") if p.is_file()] if folder.exists() else []


def _unsaved_edit(src: Path):
    from core.editing.edit_model import EditableDocument
    from core.editing.geometry import Rect

    doc = open_document(src, password=OWNER)
    edit = EditableDocument(doc)
    edit.create_text_box(0, Rect(72, 300, 260, 40), text="UNSAVED", auto_match=False)
    return doc, edit


def test_export_copy_of_a_protected_pdf_is_marked_and_removed(tmp_path, temp_dir):
    from core.printing.ops import PROTECTED_TEMP_PREFIX, prepare_current_source

    doc, edit = _unsaved_edit(_protected(tmp_path / "in.pdf"))
    try:
        source = prepare_current_source(doc, edit)
        try:
            files = _leftovers(temp_dir)
            assert len(files) == 1 and files[0].name.startswith(PROTECTED_TEMP_PREFIX)
        finally:
            source.close()
        assert _leftovers(temp_dir) == []
    finally:
        doc.close()


def test_a_failed_bake_leaves_no_temporary_copy(tmp_path, temp_dir, monkeypatch):
    import engines.bake_edits as bake
    from core.printing.ops import prepare_current_source

    doc, edit = _unsaved_edit(_protected(tmp_path / "in.pdf"))

    def failing(edit_, destination, **k):
        Path(destination).write_bytes(b"%PDF-1.7 partial " + SECRET.encode())
        raise OSError("disk full")

    monkeypatch.setattr(bake, "bake_via_backend", failing)
    try:
        with pytest.raises(OSError):
            prepare_current_source(doc, edit)
        assert _leftovers(temp_dir) == []
    finally:
        doc.close()


def test_a_failed_reopen_leaves_no_temporary_copy(tmp_path, temp_dir, monkeypatch):
    import core.printing.ops as ops

    doc, edit = _unsaved_edit(_protected(tmp_path / "in.pdf"))

    def broken(*a, **k):
        raise RuntimeError("cannot open")

    monkeypatch.setattr(ops, "open_document", broken)
    try:
        with pytest.raises(RuntimeError):
            ops.prepare_current_source(doc, edit)
        assert _leftovers(temp_dir) == []
    finally:
        doc.close()


def test_start_up_removes_copies_left_by_a_crash_at_once(tmp_path, monkeypatch):
    import core.organize.temputil as tu

    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setattr(tu, "TEMP_DIR", temp)
    monkeypatch.setattr(tu, "ORGANIZE_TEMP_DIR", temp / "organize")
    monkeypatch.setattr(tu, "ORGANIZE_CACHE_DIR", tmp_path / "cache")
    monkeypatch.setattr(tu, "RUNTIME_DIR", tmp_path)
    monkeypatch.setattr(tu, "ensure_runtime_layout", lambda: None)
    crashed = temp / "print-bake-p-0123456789abcdef.pdf"
    crashed.write_bytes(b"%PDF decrypted copy")
    recent_plain = temp / "print-bake-fedcba9876543210.pdf"
    recent_plain.write_bytes(b"%PDF")
    tu.cleanup_abandoned_temps()
    assert not crashed.exists()  # protected source: removed regardless of age
    assert recent_plain.exists()  # unprotected and recent: kept until it is a day old


def test_no_unencrypted_copy_of_the_content_remains_after_exports(tmp_path, temp_dir):
    """After a conversion from a protected PDF, nothing under the temporary folder holds its text."""
    from core.convert.pdf_to_text import pdf_to_plain_text
    from core.printing.ops import prepare_current_source

    doc, edit = _unsaved_edit(_protected(tmp_path / "in.pdf"))
    try:
        source = prepare_current_source(doc, edit)
        try:
            text = pdf_to_plain_text(source.document, None)
        finally:
            source.close()
    finally:
        doc.close()
    assert SECRET in text  # the export itself contains the text the user asked for
    for f in _leftovers(temp_dir):
        assert SECRET.encode() not in f.read_bytes(), f


def test_converted_office_files_carry_no_pdf_security(tmp_path):
    """PDF encryption applies to PDFs only: a Word file made (with the permissions password) from a
    protected PDF is an ordinary, unprotected .docx — documented in PROTECTED_SIGNED_WATERMARKED_PDFS.md."""
    pytest.importorskip("docx")
    from docx import Document

    from core.convert.pdf_to_docx import pdf_to_docx
    from core.security.protection import unlock_with_owner_password

    src = _protected(tmp_path / "in.pdf")
    doc = open_document(src, password=USER)
    try:
        assert unlock_with_owner_password(doc, OWNER)
        out = pdf_to_docx(doc, tmp_path / "out.docx", engine="text")
    finally:
        doc.close()
    text = " ".join(p.text for p in Document(str(out)).paragraphs)
    assert SECRET in text
    doc_text = (Path(__file__).resolve().parents[2] / "docs" / "PROTECTED_SIGNED_WATERMARKED_PDFS.md").read_text(
        encoding="utf-8")
    assert "not password-protected" in doc_text or "not protected" in doc_text
