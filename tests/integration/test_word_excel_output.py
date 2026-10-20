"""PDF → Excel and PDF → Word: what actually lands in the files (content, structure, numbers, Japanese)."""

from __future__ import annotations

import re
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from core.convert.pdf_to_docx import pdf_to_docx
from core.convert.tables_to_xlsx import cell_value, tables_to_xlsx
from core.document import open_document
from core.tables.detect import detect_tables

JP_FONT = fitz.Font("cjk")
A4 = (595.276, 841.89)


def _table_pdf(path: Path, *, kind: str, rows, widths=(110, 90, 90, 110)) -> Path:
    d = fitz.open()
    p = d.new_page(width=A4[0], height=A4[1])
    p.insert_font(fontname="jpf", fontbuffer=JP_FONT.buffer)
    x0, y0, rh = 60, 120, 26
    xs = [x0]
    for w in widths:
        xs.append(xs[-1] + w)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            if kind == "rects":
                p.draw_rect(fitz.Rect(xs[c], y0 + r * rh, xs[c + 1], y0 + (r + 1) * rh), color=(0, 0, 0), width=0.8)
            p.insert_text((xs[c] + 6, y0 + r * rh + 17), text, fontsize=10, fontname="jpf")
    if kind == "lines":
        for r in range(len(rows) + 1):
            p.draw_line(fitz.Point(xs[0], y0 + r * rh), fitz.Point(xs[-1], y0 + r * rh), width=0.8)
        for x in xs:
            p.draw_line(fitz.Point(x, y0), fitz.Point(x, y0 + len(rows) * rh), width=0.8)
    if kind == "thin_bars":  # borders drawn as very thin filled rectangles
        for r in range(len(rows) + 1):
            p.draw_rect(fitz.Rect(xs[0], y0 + r * rh - 0.4, xs[-1], y0 + r * rh + 0.4), color=None, fill=(0, 0, 0))
        for x in xs:
            p.draw_rect(fitz.Rect(x - 0.4, y0, x + 0.4, y0 + len(rows) * rh), color=None, fill=(0, 0, 0))
    d.save(path.as_posix())
    d.close()
    return path


ROWS = [
    ["Item", "Qty", "Price", "備考"],
    ["Apple", "12", "1,234.50", "果物"],
    ["Pear", "007", "0.25", "コード"],
    ["Total", "19", "-3", "合計"],
]


@pytest.mark.parametrize("kind", ["lines", "rects", "thin_bars"])
def test_excel_gets_the_cells_in_the_right_places_with_real_numbers(tmp_path: Path, kind: str):
    pdf = open_document(_table_pdf(tmp_path / "t.pdf", kind=kind, rows=ROWS))
    out = tmp_path / "t.xlsx"
    try:
        tables_to_xlsx(pdf, out)
    finally:
        pdf.close()
    wb = load_workbook(out.as_posix())
    ws = wb[wb.sheetnames[1]]
    assert (ws.max_row, ws.max_column) == (4, 4)
    got = [[c.value for c in row] for row in ws.iter_rows()]
    assert got[0] == ["Item", "Qty", "Price", "備考"]
    assert got[1] == ["Apple", 12, 1234.5, "果物"]       # numbers are numbers; "1,234.50" is 1234.5
    assert got[2] == ["Pear", "007", 0.25, "コード"]      # a leading zero keeps the text
    assert got[3] == ["Total", 19, -3, "合計"]
    assert isinstance(got[1][1], int) and isinstance(got[1][2], float)
    assert ws["B2"].data_type == "n" and ws["B3"].data_type == "s"
    assert "Plain numbers become Excel numbers" in wb["Info"]["A1"].value


def test_no_data_is_lost_or_duplicated_in_the_sheet(tmp_path: Path):
    pdf = open_document(_table_pdf(tmp_path / "t.pdf", kind="rects", rows=ROWS))
    try:
        out = tmp_path / "t.xlsx"
        tables_to_xlsx(pdf, out)
        page_words = [w[4] for w in pdf._backend._doc[0].get_text("words")]  # noqa: SLF001
    finally:
        pdf.close()
    ws = load_workbook(out.as_posix())[load_workbook(out.as_posix()).sheetnames[1]]
    cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    for w in page_words:
        try:
            token = str(float(w.replace(",", "")) if "." in w else int(w.replace(",", "")))
        except ValueError:
            token = w
        assert token in cells or w in cells, f"{w!r} missing from the sheet"


@pytest.mark.parametrize("text, expected", [("12", 12), ("-3", -3), ("0.25", 0.25), ("1,234.50", 1234.5), ("007", "007"),
                                            ("2026-10-16", "2026-10-16"), ("12%", "12%"), ("1,23", "1,23"), ("", ""),
                                            ("12345678901234567890", "12345678901234567890"), ("+5", "+5")])
def test_cell_value_conversion_rules(text, expected):
    assert cell_value(text) == expected


def test_page_without_a_table_gives_an_honest_empty_workbook(tmp_path: Path):
    d = fitz.open()
    d.new_page().insert_text((72, 100), "No table here")
    src = tmp_path / "n.pdf"
    d.save(src.as_posix())
    d.close()
    pdf = open_document(src)
    try:
        tables_to_xlsx(pdf, tmp_path / "n.xlsx")
    finally:
        pdf.close()
    assert load_workbook((tmp_path / "n.xlsx").as_posix())["Info"]["A2"].value == "No tables detected."


def test_known_false_positive_guard_on_real_samples():
    """Ordinary form / text samples must not turn into 'tables' now that rectangles count as rules."""
    samples = Path(__file__).resolve().parents[1] / "samples"
    for name in ("sample_text.pdf", "sample_invitation_form.pdf", "sample_jp_admin_form.pdf", "sample_mixed_fonts.pdf", "sample_image_page.pdf"):
        d = fitz.open((samples / name).as_posix())
        try:
            assert detect_tables(d) == [], f"{name} produced a table"
        finally:
            d.close()


# ------------------------------------------------------------------ Word
def _doc_pdf(path: Path, size=A4) -> Path:
    d = fitz.open()
    p = d.new_page(width=size[0], height=size[1])
    p.insert_font(fontname="jpf", fontbuffer=JP_FONT.buffer)
    p.insert_text((72, 90), "Quarterly Report", fontsize=22, fontname="helv")
    p.insert_text((72, 130), "First paragraph of ordinary body text that stays in order.", fontsize=11, fontname="helv")
    p.insert_text((72, 150), "Second paragraph follows directly after the first.", fontsize=11, fontname="helv")
    p.insert_text((72, 190), "日本語の段落です。東京都千代田区。", fontsize=11, fontname="jpf")
    p.insert_text((72, 230), "Mixed 注文番号 2026-17 line", fontsize=11, fontname="jpf")
    pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 60), False)
    pm.set_rect(pm.irect, (200, 40, 40))
    p.insert_image(fitz.Rect(72, 280, 232, 360), pixmap=pm)
    d.save(path.as_posix())
    d.close()
    return path


def test_word_file_has_text_in_order_japanese_image_and_page_size(tmp_path: Path):
    from docx import Document

    pdf = open_document(_doc_pdf(tmp_path / "d.pdf"))
    out = tmp_path / "d.docx"
    try:
        pdf_to_docx(pdf, out)
    finally:
        pdf.close()
    doc = Document(out.as_posix())
    text = "\n".join(p.text for p in doc.paragraphs)
    flat = re.sub(r"\s+", "", text)
    order = ["QuarterlyReport", "Firstparagraph", "Secondparagraph", "日本語の段落です", "Mixed注文番号2026-17"]
    pos = [flat.find(re.sub(r"\s+", "", t)) for t in order]
    assert all(p >= 0 for p in pos), (pos, text)
    assert pos == sorted(pos), "paragraph order changed"
    assert len(doc.inline_shapes) + len(doc.part.package.image_parts) >= 1, "image missing"
    sec = doc.sections[0]
    assert sec.page_width.pt == pytest.approx(A4[0], abs=3) and sec.page_height.pt == pytest.approx(A4[1], abs=3)
    # real, editable text: runs carry text, not a page picture
    assert sum(len(r.text) for p in doc.paragraphs for r in p.runs) > 60


def test_word_file_of_a_landscape_letter_page_keeps_orientation(tmp_path: Path):
    from docx import Document

    pdf = open_document(_doc_pdf(tmp_path / "d.pdf", size=(792.0, 612.0)))
    try:
        pdf_to_docx(pdf, tmp_path / "d.docx")
    finally:
        pdf.close()
    sec = Document((tmp_path / "d.docx").as_posix()).sections[0]
    assert sec.page_width.pt > sec.page_height.pt
    assert sec.page_width.pt == pytest.approx(792.0, abs=3)
