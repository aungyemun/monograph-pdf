"""Area OCR, word positions on turned pages and sideways scans, the searchable layer's placement, no
duplicate OCR layer, OCR of the current (reordered) document, and text / Word export."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest

from core.document import open_document
from core.editing.geometry import Rect
from core.ocr.engine import get_ocr_engine
from core.ocr.paths import ocr_ready
from core.ocr.pipeline import OcrOptions, has_invisible_text, run_ocr
from core.ocr.text_export import document_text, join_words, page_text, write_docx, write_text


@pytest.fixture(scope="module")
def engine():
    ok, detail = ocr_ready()
    if not ok:
        pytest.skip(detail)
    return get_ocr_engine(force_reload=True)


def _scan(path: Path, lines: list[str], *, width=600, height=300, rotate_page=0, sideways=False, turn=0) -> Path:
    """An image-only page with ``lines`` of text; optionally a /Rotate page or a scan lying sideways
    (``turn``: the stored picture is turned by this angle)."""
    src = fitz.open()
    p = src.new_page(width=width, height=height)
    y = 60
    for line in lines:
        p.insert_text((40, y), line, fontsize=30)
        y += 60
    matrix = fitz.Matrix(200 / 72, 200 / 72).prerotate(turn or (90 if sideways else 0))
    pix = p.get_pixmap(matrix=matrix, alpha=False)
    src.close()
    w, h = pix.width * 72 / 200, pix.height * 72 / 200
    doc = fitz.open()
    page = doc.new_page(width=w, height=h)
    page.insert_image(page.rect, pixmap=pix)
    if rotate_page:
        page.set_rotation(rotate_page)
    doc.save(path)
    doc.close()
    return path


def _words(result):
    return [w for pr in result.pages for w in pr.words]


def test_area_ocr_reads_only_the_selected_area(engine, tmp_path):
    pdf = open_document(_scan(tmp_path / "s.pdf", ["FIRST LINE HERE", "SECOND LINE THERE"]))
    try:
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], region=Rect(20, 80, 560, 60),
                                       make_searchable=False), engine=engine)
    finally:
        pdf.close()
    text = " ".join(w.text for w in _words(res))
    assert "SECOND" in text and "FIRST" not in text
    for w in _words(res):  # positions are page positions, not positions inside the area picture
        assert 80 <= w.rect.y + w.rect.h / 2 <= 140


def test_area_outside_the_page_is_reported_not_crashed(engine, tmp_path):
    pdf = open_document(_scan(tmp_path / "s.pdf", ["TEXT"]))
    try:
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], region=Rect(5000, 5000, 100, 100),
                                       make_searchable=False), engine=engine)
    finally:
        pdf.close()
    assert res.pages[0].skipped and "outside" in res.pages[0].skip_reason


@pytest.mark.parametrize("turn,rotate_page", [(90, 270), (270, 90), (180, 180)])
def test_turned_pages_are_read_upright_and_the_layer_lines_up(engine, tmp_path, turn, rotate_page):
    """A scan stored turned and shown upright through /Rotate (the usual case for rotated pages):
    words are found in display space, and the searchable text sits on the words as shown."""
    src = _scan(tmp_path / "r.pdf", ["ROTATED PAGE WORDS"], turn=turn, rotate_page=rotate_page)
    out = tmp_path / "r_ocr.pdf"
    pdf = open_document(src)
    try:
        info = pdf.page(0)
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], skip_native_text=False, make_searchable=True,
                                       apply_orientation=False), engine=engine, destination=out)
        disp_w, disp_h = info.display_width_pt, info.display_height_pt
    finally:
        pdf.close()
    words = {w.text: w.rect for w in _words(res)}
    assert "ROTATED" in words
    for r in words.values():
        assert -1 <= r.x and r.x1 <= disp_w + 1 and -1 <= r.y and r.y1 <= disp_h + 1
    assert disp_w > disp_h  # shown as a landscape line of text
    from core.page_space import rect_to_display

    with fitz.open(out) as d:
        page = d[0]
        placed = {w[4]: rect_to_display(page, fitz.Rect(w[:4])) for w in page.get_text("words")}
    q = words["ROTATED"]
    assert placed["ROTATED"].contains(fitz.Point(q.x + q.w / 2, q.y + q.h / 2))


def test_searchable_layer_lines_up_on_a_sideways_scan(engine, tmp_path):
    """The scan lies sideways: orientation detection turns it for recognition, and the invisible text
    is written along the picture's text so selecting it selects the right place."""
    src = _scan(tmp_path / "side.pdf", ["SIDEWAYS SCAN TEXT LINE", "ANOTHER LINE OF WORDS"], sideways=True)
    out = tmp_path / "side_ocr.pdf"
    pdf = open_document(src)
    try:
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], skip_native_text=False, make_searchable=True),
                      engine=engine, destination=out)
    finally:
        pdf.close()
    if not res.pages[0].rotation_applied:
        pytest.skip("orientation detection did not report the turn on this engine build")
    ocr_words = {w.text: w.rect for w in _words(res)}
    with fitz.open(out) as d:
        placed = {w[4]: fitz.Rect(w[:4]) for w in d[0].get_text("words")}
    assert "SIDEWAYS" in placed
    r, q = placed["SIDEWAYS"], ocr_words["SIDEWAYS"]
    centre_ocr = ((q.x + q.x1) / 2, (q.y + q.y1) / 2)
    assert r.contains(fitz.Point(centre_ocr)) or abs((r.x0 + r.x1) / 2 - centre_ocr[0]) < 12


def test_no_second_text_layer_and_ocr_uses_the_current_page_order(engine, tmp_path):
    src = tmp_path / "two.pdf"
    a = _scan(tmp_path / "a.pdf", ["PAGE ALPHA"])
    b = _scan(tmp_path / "b.pdf", ["PAGE BRAVO"])
    with fitz.open(a) as da, fitz.open(b) as db:
        da.insert_pdf(db)
        da.save(src)
    first = tmp_path / "first.pdf"
    pdf = open_document(src)
    try:
        doc = pdf._backend._doc  # noqa: SLF001 — reorder the open document (unsaved)
        doc.move_page(1, 0)
        pdf.invalidate_pages()
        run_ocr(pdf, OcrOptions(languages="eng", pages=[0, 1], skip_native_text=False), engine=engine, destination=first)
    finally:
        pdf.close()
    with fitz.open(first) as d:
        assert "BRAVO" in d[0].get_text() and "ALPHA" in d[1].get_text()  # the open order, not the file's
        assert has_invisible_text(d[0])
    again = tmp_path / "again.pdf"
    pdf = open_document(first)
    try:
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], skip_native_text=False), engine=engine, destination=again)
    finally:
        pdf.close()
    assert "already has a searchable text layer" in res.pages[0].skip_reason
    with fitz.open(again) as d:
        assert d[0].get_text().count("BRAVO") == 1


def test_text_and_word_export(engine, tmp_path):
    from docx import Document

    src = _scan(tmp_path / "doc.pdf", ["Quarterly Report", "first paragraph line one", "first paragraph line two"])
    pdf = open_document(src)
    try:
        res = run_ocr(pdf, OcrOptions(languages="eng", pages=[0], skip_native_text=False, make_searchable=False),
                      engine=engine)
    finally:
        pdf.close()
    text = document_text(res.pages)
    assert text.startswith("--- Page 1 ---") and "Quarterly" in text
    txt = write_text(text, tmp_path / "out.txt")
    assert txt.read_bytes().startswith(b"\xef\xbb\xbf") and "Quarterly" in txt.read_text(encoding="utf-8-sig")
    docx = write_docx(res.pages, tmp_path / "out.docx", title="Scan")
    paras = [p.text for p in Document(str(docx)).paragraphs if p.text.strip()]
    assert any("Quarterly" in p for p in paras)
    edited = write_docx(res.pages, tmp_path / "edited.docx", edited_text="--- Page 1 ---\nCorrected title\n\nBody text")
    assert [p.text for p in Document(str(edited)).paragraphs if p.text.strip()] == ["Corrected title", "Body text"]


def test_japanese_words_are_joined_without_spaces():
    assert join_words(["日本", "語", "の", "OCR", "text"]) == "日本語の OCR text"

    class W:
        def __init__(self, text, block, par, line, h=10.0):
            self.text, self.block, self.par, self.line = text, block, par, line
            self.rect = Rect(0, 0, 10, h)

    words = [W("見出し", 1, 1, 1, 20), W("本文", 2, 1, 1), W("です", 2, 1, 1), W("次", 2, 1, 2)]
    assert page_text(words) == "見出し\n\n本文です\n次"
