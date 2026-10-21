"""PDF -> PowerPoint slide geometry (A4 portrait, A3 landscape, Legal, custom, mixed sizes with every size
policy) and Office document properties (mapped from the PDF; no library template values)."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import fitz
import pytest

from core.convert.office_metadata import parse_pdf_date
from core.convert.pdf_to_pptx import EMU_PER_PT, PptxReport, page_groups, pdf_to_pptx, pdf_to_pptx_split
from core.document import open_document

A4 = (595.28, 841.89)
A3L = (1190.55, 841.89)
LEGAL = (612.0, 1008.0)
CUSTOM = (400.0, 250.0)
TEMPLATE_VALUES = ("python-docx", "python-pptx", "Steve Canny", "openpyxl", "generated using")
META = {"title": "Cell atlas", "author": "A. Researcher", "subject": "Microscopy", "keywords": "aging; RNA-seq",
        "creationDate": "D:20240102030405+09'00'", "modDate": "D:20250607080910Z"}


def _pdf(path: Path, sizes, metadata=META) -> Path:
    doc = fitz.open()
    for i, (w, h) in enumerate(sizes):
        page = doc.new_page(width=w, height=h)
        page.insert_text((30, 50), f"Page {i + 1}", fontsize=16)
        page.draw_rect(page.rect + (5, 5, -5, -5), color=(1, 0, 0), width=2)
    if metadata:
        doc.set_metadata(metadata)
    doc.save(path)
    doc.close()
    return path


def _slides(path: Path):
    from pptx import Presentation

    prs = Presentation(path.as_posix())
    return prs, [(round(prs.slide_width / EMU_PER_PT), round(prs.slide_height / EMU_PER_PT))]


def _pictures(path: Path):
    from pptx import Presentation

    prs = Presentation(path.as_posix())
    out = []
    for slide in prs.slides:
        pics = [s for s in slide.shapes if s.shape_type == 13]
        out.append([(s.left, s.top, s.width, s.height) for s in pics])
    return prs, out


@pytest.mark.parametrize("size", [A4, A3L, LEGAL, CUSTOM], ids=["A4-portrait", "A3-landscape", "Legal", "custom"])
def test_single_size_gives_that_exact_slide_size(tmp_path, size):
    pdf = open_document(_pdf(tmp_path / "in.pdf", [size, size]))
    try:
        report = PptxReport()
        out = pdf_to_pptx(pdf, tmp_path / "out.pptx", report=report)
    finally:
        pdf.close()
    prs, ((w, h),) = _slides(out)
    assert abs(w - size[0]) <= 1 and abs(h - size[1]) <= 1
    assert report.strategy == "uniform" and len(prs.slides) == 2
    _prs, pics = _pictures(out)
    for slide in pics:  # each page fills its slide exactly, proportions kept
        (left, top, pw, ph), = slide
        assert (left, top) == (0, 0) and abs(pw - prs.slide_width) <= 2 and abs(ph - prs.slide_height) <= 2


def _mixed(tmp_path):
    return _pdf(tmp_path / "mixed.pdf", [A4, A3L, LEGAL])


def test_mixed_sizes_true_size_and_fit_policies(tmp_path):
    src = _mixed(tmp_path)
    for policy in ("true_size", "fit"):
        pdf = open_document(src)
        try:
            out = pdf_to_pptx(pdf, tmp_path / f"{policy}.pptx", size_policy=policy)
        finally:
            pdf.close()
        prs, pics = _pictures(out)
        sw, sh = prs.slide_width, prs.slide_height
        assert abs(sw / EMU_PER_PT - A3L[0]) <= 1 and abs(sh / EMU_PER_PT - LEGAL[1]) <= 1  # widest × tallest
        for (left, top, w, h), (pw, ph) in zip([p[0] for p in pics], (A4, A3L, LEGAL)):
            assert abs(w / h - pw / ph) < 0.01  # proportions preserved, never stretched
            assert left >= 0 and top >= 0 and left + w <= sw + 2 and top + h <= sh + 2  # never cropped
            if policy == "true_size":
                assert abs(w / EMU_PER_PT - pw) <= 1  # true scale
            else:
                assert abs(w - sw) <= 2 or abs(h - sh) <= 2  # fills one dimension


def test_first_page_and_custom_policies(tmp_path):
    src = _mixed(tmp_path)
    pdf = open_document(src)
    try:
        r1 = PptxReport()
        first = pdf_to_pptx(pdf, tmp_path / "first.pptx", size_policy="first_page", report=r1)
        r2 = PptxReport()
        custom = pdf_to_pptx(pdf, tmp_path / "custom.pptx", size_policy="custom", custom_size_pt=(960.0, 540.0),
                             report=r2)
        with pytest.raises(ValueError):
            pdf_to_pptx(pdf, tmp_path / "x.pptx", size_policy="custom")
    finally:
        pdf.close()
    prs, ((w, h),) = _slides(first)
    assert abs(w - A4[0]) <= 1 and abs(h - A4[1]) <= 1 and r1.strategy == "first page size"
    prs, ((w, h),) = _slides(custom)
    assert (w, h) == (960, 540) and r2.strategy == "custom slide size"
    _prs, pics = _pictures(custom)
    for (left, top, pw, ph), (ow, oh) in zip([p[0] for p in pics], (A4, A3L, LEGAL)):
        assert abs(pw / ph - ow / oh) < 0.01 and left + pw <= prs.slide_width + 2 and top + ph <= prs.slide_height + 2


def test_split_policy_writes_one_presentation_per_size(tmp_path):
    src = _pdf(tmp_path / "m.pdf", [A4, A3L, A4, LEGAL])
    pdf = open_document(src)
    try:
        groups = page_groups(pdf)
        outs = pdf_to_pptx_split(pdf, tmp_path / "deck.pptx")
    finally:
        pdf.close()
    assert [len(m) for _l, m in groups] == [2, 1, 1] and groups[0][1] == [0, 2]
    assert len(outs) == 3 and all(o.exists() for o in outs)
    assert "A4" in outs[0].name and outs[0].name.startswith("deck (")
    sizes = [_slides(o)[1][0] for o in outs]
    assert abs(sizes[0][0] - 595) <= 1 and abs(sizes[1][0] - 1191) <= 1 and abs(sizes[2][1] - 1008) <= 1
    from pptx import Presentation

    assert [len(Presentation(o.as_posix()).slides) for o in outs] == [2, 1, 1]


def _core_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("docProps/core.xml").decode("utf-8")


def test_office_outputs_carry_pdf_metadata_and_no_template_values(tmp_path):
    from core.convert.pdf_to_docx import pdf_to_docx
    from core.convert.tables_to_xlsx import tables_to_xlsx

    src = _pdf(tmp_path / "meta.pdf", [A4])
    pdf = open_document(src)
    try:
        outs = [pdf_to_pptx(pdf, tmp_path / "m.pptx"), pdf_to_pptx(pdf, tmp_path / "e.pptx", mode="editable"),
                pdf_to_docx(pdf, tmp_path / "t.docx", engine="text"), tables_to_xlsx(pdf, tmp_path / "x.xlsx")]
        try:
            outs.append(pdf_to_docx(pdf, tmp_path / "l.docx", engine="layout"))
        except Exception:  # noqa: BLE001 — layout engine optional
            pass
    finally:
        pdf.close()
    import getpass

    user = getpass.getuser()
    for out in outs:
        xml = _core_xml(out)
        for bad in TEMPLATE_VALUES:
            assert bad not in xml, (out.name, bad)
        if len(user) > 3:
            assert f">{user}<" not in xml, (out.name, "local user name")
        assert "Cell atlas" in xml and "A. Researcher" in xml, out.name
        assert "2024-01-01T18:04:05Z" in xml, out.name  # 03:04:05 +09:00 → UTC
    assert "Microscopy" in _core_xml(outs[0]) and "aging; RNA-seq" in _core_xml(outs[0])


def test_missing_metadata_is_left_empty_not_invented(tmp_path):
    src = _pdf(tmp_path / "bare.pdf", [A4], metadata={"title": "", "author": "", "creator": "", "producer": ""})
    pdf = open_document(src)
    try:
        out = pdf_to_pptx(pdf, tmp_path / "bare.pptx")
    finally:
        pdf.close()
    from pptx import Presentation

    cp = Presentation(out.as_posix()).core_properties
    assert cp.title == "" and cp.author == "" and cp.comments == "" and cp.last_modified_by == ""


def test_parse_pdf_date():
    assert parse_pdf_date("D:20261021093000+09'00'").isoformat() == "2026-10-21T00:30:00+00:00"
    assert parse_pdf_date("D:2026").year == 2026
    assert parse_pdf_date("") is None and parse_pdf_date("garbage") is None


@pytest.mark.parametrize("engine", ["text", "layout"])
def test_pdf_to_docx_sections_follow_every_page_size_and_orientation(tmp_path, engine):
    from docx import Document
    from docx.enum.section import WD_ORIENT

    from core.convert.pdf_to_docx import layout_engine_available, pdf_to_docx

    if engine == "layout" and not layout_engine_available():
        pytest.skip("pdf2docx not installed")
    sizes = [A4, A3L, (612.0, 792.0), LEGAL, CUSTOM]
    pdf = open_document(_pdf(tmp_path / "m.pdf", sizes))
    try:
        out = pdf_to_docx(pdf, tmp_path / f"{engine}.docx", engine=engine)
    finally:
        pdf.close()
    secs = Document(out.as_posix()).sections
    assert len(secs) == len(sizes)
    for s, (w, h) in zip(secs, sizes):
        assert abs(s.page_width.pt - w) <= 1 and abs(s.page_height.pt - h) <= 1
        assert s.orientation == (WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT)
    assert not list(tmp_path.glob("*.part"))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    text = " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    assert all(f"Page {i + 1}" in text.replace("  ", " ") for i in range(len(sizes))), text[:300]
