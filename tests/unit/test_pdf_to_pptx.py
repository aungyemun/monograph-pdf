"""PDF → PowerPoint: slide size, mixed pages, rotated pages, overflow, text boxes, picture fidelity."""

from __future__ import annotations

import io
from pathlib import Path

import fitz
import pytest
from PIL import Image, ImageChops
from pptx import Presentation

from core.convert.pdf_to_pptx import EMU_PER_MM, EMU_PER_PT, PptxReport, pdf_to_pptx, powerpoint_family
from core.document import open_document

A4 = (595.276, 841.89)
LETTER = (612.0, 792.0)


def _pdf(path: Path, sizes, *, rotate=None, text=True, table=False, image=False) -> Path:
    doc = fitz.open()
    for n, (w, h) in enumerate(sizes):
        page = doc.new_page(width=w, height=h)
        if text:
            page.insert_text((50, 80), f"Page {n + 1} heading", fontsize=22, fontname="helv", color=(0.8, 0.1, 0.1))
            page.insert_text((50, 120), "Body text in Times 12 pt with several words.", fontsize=12, fontname="tiro")
            page.insert_text((50, 140), "Courier line 10 pt", fontsize=10, fontname="cour")
            for k in range(25):
                page.insert_text((50, 170 + k * 14), f"Line {k} of a text heavy page lorem ipsum dolor sit amet", fontsize=9)
        if table:
            for r in range(5):
                for c in range(4):
                    page.draw_rect(fitz.Rect(50 + c * 80, 560 + r * 24, 130 + c * 80, 584 + r * 24), color=(0, 0, 0), width=0.8)
                    page.insert_text((56 + c * 80, 576 + r * 24), f"R{r}C{c}", fontsize=9)
        if image:
            pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 40), False)
            pm.set_rect(pm.irect, (30, 120, 220))
            page.insert_image(fitz.Rect(300, 40, 420, 120), pixmap=pm)
        if rotate and rotate[n]:
            page.set_rotation(rotate[n])
    doc.save(path.as_posix())
    doc.close()
    return path


def _convert(tmp_path: Path, pdf: Path, **kw):
    doc = open_document(pdf)
    rep = PptxReport()
    out = tmp_path / "out.pptx"
    try:
        pdf_to_pptx(doc, out, report=rep, **kw)
    finally:
        doc.close()
    return Presentation(out.as_posix()), rep


def _mm(emu: int) -> float:
    return emu / EMU_PER_MM


@pytest.mark.parametrize(
    "size, expect_mm, label",
    [
        (A4, (210.0, 297.0), "A4 portrait"),
        ((A4[1], A4[0]), (297.0, 210.0), "A4 landscape"),
        (LETTER, (215.9, 279.4), "Letter portrait"),
        ((419.53, 595.28), (148.0, 210.0), "A5 portrait"),
    ],
)
@pytest.mark.parametrize("mode", ["layout", "editable"])
def test_standard_sheets_keep_their_size(tmp_path: Path, size, expect_mm, label, mode):
    prs, rep = _convert(tmp_path, _pdf(tmp_path / "a.pdf", [size]), mode=mode)
    assert (_mm(prs.slide_width), _mm(prs.slide_height)) == pytest.approx(expect_mm, abs=0.01)
    assert rep.slide_label == label
    assert rep.strategy == "uniform"
    # never PowerPoint's 16:9 default
    assert prs.slide_width / prs.slide_height != pytest.approx(16 / 9, rel=0.01)


def test_nonstandard_page_uses_exact_point_size(tmp_path: Path):
    prs, rep = _convert(tmp_path, _pdf(tmp_path / "a.pdf", [(400.0, 333.0)]))
    assert prs.slide_width == 400 * EMU_PER_PT and prs.slide_height == 333 * EMU_PER_PT
    assert rep.slide_label == ""


@pytest.mark.parametrize("rot, expect", [(0, "A4 portrait"), (90, "A4 landscape"), (180, "A4 portrait"), (270, "A4 landscape")])
def test_rotated_pages_use_displayed_size(tmp_path: Path, rot: int, expect: str):
    prs, rep = _convert(tmp_path, _pdf(tmp_path / "a.pdf", [A4], rotate=[rot]))
    assert rep.slide_label == expect
    landscape = rot in (90, 270)
    assert (prs.slide_width > prs.slide_height) == landscape


def test_mixed_sizes_use_bounding_sheet_without_stretching(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4, (A4[1], A4[0]), LETTER])
    prs, rep = _convert(tmp_path, pdf)
    assert rep.strategy == "bounding sheet" and rep.notes
    assert prs.slide_width >= 297 * EMU_PER_MM - 5000 and prs.slide_height >= 297 * EMU_PER_MM - 5000
    for slide, (w, h) in zip(prs.slides, [A4, (A4[1], A4[0]), LETTER]):
        pic = next(s for s in slide.shapes if s.shape_type == 13)
        # true aspect ratio, inside the slide, centered
        assert pic.width / pic.height == pytest.approx(w / h, rel=0.005)
        assert pic.left >= 0 and pic.top >= 0
        assert pic.left + pic.width <= prs.slide_width and pic.top + pic.height <= prs.slide_height
        assert abs((pic.left + pic.width / 2) - prs.slide_width / 2) < 20000
        assert abs((pic.top + pic.height / 2) - prs.slide_height / 2) < 20000


@pytest.mark.parametrize("mode", ["layout", "editable"])
def test_nothing_extends_outside_the_slide(tmp_path: Path, mode: str):
    pdf = _pdf(tmp_path / "a.pdf", [A4, LETTER, A4], text=True, table=True, image=True)
    prs, rep = _convert(tmp_path, pdf, mode=mode)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (mode, shape.shape_type)
            assert shape.left + shape.width <= prs.slide_width + 1, (mode, shape.name)
            assert shape.top + shape.height <= prs.slide_height + 1, (mode, shape.name)
    assert len(prs.slides) == 3


def test_layout_mode_picture_matches_the_pdf_page(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4], table=True, image=True)
    prs, _ = _convert(tmp_path, pdf, mode="layout")
    slide = prs.slides[0]
    pics = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pics) == 1 and len(slide.shapes) == 1
    pic = pics[0]
    assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, prs.slide_width, prs.slide_height)
    img = Image.open(io.BytesIO(pic.image.blob)).convert("L")
    ref = fitz.open(pdf.as_posix())[0].get_pixmap(matrix=fitz.Matrix(img.width / A4[0], img.height / A4[1]), alpha=False)
    expected = Image.frombytes("RGB", (ref.width, ref.height), ref.samples).convert("L").resize(img.size)
    diff = ImageChops.difference(img, expected).resize((100, 140))
    assert sum(diff.getdata()) / (100 * 140) < 3.0


def test_layout_mode_rotated_page_picture_is_upright(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4], rotate=[90])
    prs, _ = _convert(tmp_path, pdf)
    pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    img = Image.open(io.BytesIO(pic.image.blob))
    assert img.width > img.height
    assert img.width / img.height == pytest.approx(A4[1] / A4[0], rel=0.01)


def test_editable_mode_text_boxes_sit_where_the_text_is(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4], text=True)
    prs, rep = _convert(tmp_path, pdf, mode="editable")
    assert rep.text_boxes > 20
    page = fitz.open(pdf.as_posix())[0]
    boxes = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    heading = next(b for b in boxes if b.text_frame.text.startswith("Page 1 heading"))
    w = next(w for w in page.get_text("words") if w[4] == "Page")
    # same position within a point (EMU → pt) and a font size equal to the PDF's
    assert abs(heading.left / EMU_PER_PT - w[0]) < 1.5
    assert abs(heading.top / EMU_PER_PT - w[1]) < 4.0
    run = heading.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(22.0, abs=0.2)
    assert run.font.name == "Arial"  # Helvetica → Arial
    assert str(run.font.color.rgb) == "CC1A1A"
    # text really is text (searchable), not only a picture
    assert any("Body text in Times" in b.text_frame.text for b in boxes)
    # the picture behind the text does not contain the text again
    bg = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    img = Image.open(io.BytesIO(bg.image.blob)).convert("L")
    dark = sum(1 for p in img.getdata() if p < 100)
    assert dark < img.width * img.height * 0.002


def test_editable_mode_on_rotated_page_turns_the_text_with_the_page(tmp_path: Path):
    """The text of a /Rotate 90 page is horizontal in the file but shown turned a quarter clockwise:
    the slide is landscape and every text box carries that 90° angle, inside the slide."""
    pdf = _pdf(tmp_path / "a.pdf", [A4], rotate=[90], text=True)
    prs, rep = _convert(tmp_path, pdf, mode="editable")
    assert prs.slide_width > prs.slide_height
    boxes = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert boxes and rep.rotated_text == len(boxes)
    assert all(abs(b.rotation - 90.0) < 0.6 for b in boxes)
    heading = next(b for b in boxes if b.text_frame.text.startswith("Page 1 heading"))
    # shown at the upper right of the landscape sheet (the PDF's top-left corner after a clockwise turn)
    cx = heading.left + heading.width / 2
    assert cx > prs.slide_width * 0.8


def test_editable_mode_keeps_rotated_text_angle(tmp_path: Path):
    doc = fitz.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((100, 400), "Turned text", fontsize=14, rotate=90)
    pdf = tmp_path / "r.pdf"
    doc.save(pdf.as_posix())
    doc.close()
    prs, rep = _convert(tmp_path, pdf, mode="editable")
    boxes = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert rep.rotated_text == 1
    assert boxes[0].rotation == pytest.approx(270.0, abs=1.0)  # PDF rotate=90 is a quarter turn anticlockwise


def test_tables_and_images_survive_in_editable_mode(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4], table=True, image=True)
    prs, rep = _convert(tmp_path, pdf, mode="editable")
    bg = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    img = Image.open(io.BytesIO(bg.image.blob)).convert("RGB")
    sx = img.width / A4[0]
    # table rule pixels (black) along the first row's top edge and the blue image block are still drawn
    px = img.load()
    assert any(px[int((50 + c) * sx), int(560 * (img.height / A4[1]))][0] < 80 for c in range(0, 300, 7))
    blue = px[int(360 * sx), int(80 * (img.height / A4[1]))]
    assert blue[2] > 180 and blue[0] < 80
    assert any(s.has_text_frame and "R0C0" in s.text_frame.text for s in prs.slides[0].shapes)


@pytest.mark.parametrize(
    "pdf_font, family",
    [
        ("ABCDEF+Arial-BoldMT", "Arial"),
        ("Helvetica-Oblique", "Arial"),
        ("TimesNewRomanPSMT", "Times New Roman"),
        ("Times-Roman", "Times New Roman"),
        ("CourierNewPS-BoldMT", "Courier New"),
        ("XYZABC+SegoeUI", "Segoe UI"),
        ("Calibri", "Calibri"),
    ],
)
def test_pdf_font_names_map_to_powerpoint_families(pdf_font: str, family: str):
    assert powerpoint_family(pdf_font) == family


def test_unknown_mode_is_rejected(tmp_path: Path):
    doc = open_document(_pdf(tmp_path / "a.pdf", [A4]))
    try:
        with pytest.raises(ValueError):
            pdf_to_pptx(doc, tmp_path / "x.pptx", mode="pixel-perfect")
    finally:
        doc.close()


def test_page_selection_and_reopen_in_python_pptx(tmp_path: Path):
    pdf = _pdf(tmp_path / "a.pdf", [A4, A4, A4])
    prs, rep = _convert(tmp_path, pdf, pages=[0, 2])
    assert len(prs.slides) == 2 and rep.pages == 2
    assert "A4 portrait" in rep.summary()


def test_rotated_text_stays_inside_the_slide_when_turned(tmp_path: Path):
    import math

    doc = fitz.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((20, 800), "Very close to the left edge", fontsize=14, rotate=90)
    pdf = tmp_path / "r.pdf"
    doc.save(pdf.as_posix())
    doc.close()
    prs, _ = _convert(tmp_path, pdf, mode="editable")
    for b in (s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()):
        rad = math.radians(b.rotation)
        hw = (abs(b.width * math.cos(rad)) + abs(b.height * math.sin(rad))) / 2
        hh = (abs(b.width * math.sin(rad)) + abs(b.height * math.cos(rad))) / 2
        cx, cy = b.left + b.width / 2, b.top + b.height / 2
        assert cx - hw >= -1 and cx + hw <= prs.slide_width + 1
        assert cy - hh >= -1 and cy + hh <= prs.slide_height + 1
