"""Offline convert/export path tests — artifacts under tests/outputs."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest

from core.convert.capabilities import (
    Fidelity,
    UnsupportedConversionError,
    list_capabilities,
)
from core.convert.copy_pdf import copy_pdf_preserving_metadata
from core.convert.extract_images import extract_embedded_images
from core.convert.flatten_export import export_flattened_pdf
from core.convert.images_to_pdf import images_to_pdf
from core.convert.office_to_pdf import office_to_pdf, office_to_pdf_supported
from core.convert.pages_to_pdf import pages_to_pdf
from core.convert.pdf_to_docx import FIDELITY_NOTICE as DOCX_NOTICE
from core.convert.pdf_to_docx import pdf_to_docx
from core.convert.pdf_to_images import pdf_pages_to_images
from core.convert.pdf_to_pptx import FIDELITY_NOTICE as PPTX_NOTICE
from core.convert.pdf_to_pptx import pdf_to_pptx
from core.convert.pdf_to_text import pdf_to_plain_text
from core.convert.tables_to_xlsx import FIDELITY_NOTICE as XLSX_NOTICE
from core.convert.tables_to_xlsx import tables_to_xlsx
from core.convert.watermark import apply_headers_footers, apply_page_numbers, apply_watermark
from core.document import open_document
from core.editing.edit_model import EditableDocument
from core.editing.geometry import Rect
from core.editing.image_box import ImageBox, export_image_box
from engines.bake_edits import bake_into_document

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
OUT = Path(__file__).resolve().parents[1] / "outputs" / "generated"
IMG_DIR = OUT / "convert_images"


@pytest.fixture(scope="module", autouse=True)
def _fixtures():
    from tests.fixtures.make_samples import main

    main()
    OUT.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    # Tiny PNG for images→PDF and ImageBox
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 80, 60), 0)
    pix.set_rect(pix.irect, (40, 120, 200))
    pix.save((IMG_DIR / "swatch.png").as_posix())


def test_capabilities_registry():
    caps = list_capabilities()
    keys = {c.key for c in caps}
    assert "images_to_pdf" in keys
    assert "office_to_pdf" in keys
    office = next(c for c in caps if c.key == "office_to_pdf")
    # Implemented exactly when an office suite is installed on this machine
    assert office.implemented is office_to_pdf_supported()
    if not office.implemented:
        assert office.fidelity == Fidelity.UNSUPPORTED
    assert any(c.fidelity == Fidelity.BEST_EFFORT for c in caps)
    assert DOCX_NOTICE and XLSX_NOTICE and PPTX_NOTICE


def test_images_to_pdf():
    dest = OUT / "convert_images_to_pdf.pdf"
    images_to_pdf([IMG_DIR / "swatch.png", IMG_DIR / "swatch.png"], dest)
    pdf = open_document(dest)
    try:
        assert pdf.page_count == 2
    finally:
        pdf.close()


def test_pdf_to_png_jpeg_tiff():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    try:
        pngs = pdf_pages_to_images(pdf, OUT / "conv_png", pages=[0], fmt="png")
        jpgs = pdf_pages_to_images(pdf, OUT / "conv_jpg", pages=[0], fmt="jpeg")
        tifs = pdf_pages_to_images(pdf, OUT / "conv_tif", pages=[0], fmt="tiff")
        assert pngs[0].is_file() and jpgs[0].is_file() and tifs[0].is_file()
    finally:
        pdf.close()
    with pytest.raises(ValueError):
        pdf = open_document(SAMPLES / "sample_text.pdf")
        try:
            pdf_pages_to_images(pdf, OUT, fmt="gif")
        finally:
            pdf.close()


def test_extract_embedded_images():
    pdf = open_document(SAMPLES / "sample_image_page.pdf")
    try:
        out = OUT / "extracted_imgs"
        written = extract_embedded_images(pdf, out)
        assert isinstance(written, list)
    finally:
        pdf.close()


def test_pdf_to_plain_text():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    dest = OUT / "convert_plain.txt"
    try:
        text = pdf_to_plain_text(pdf, dest)
        assert "alpha-bravo-charlie" in text.lower() or "fixture" in text.lower()
        assert dest.is_file()
    finally:
        pdf.close()


def test_pages_to_pdf_and_metadata_copy_and_flatten():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    try:
        pages_to_pdf(pdf, [0, 2], OUT / "convert_pages.pdf")
        copy_pdf_preserving_metadata(pdf, OUT / "convert_meta_copy.pdf")
        export_flattened_pdf(pdf, OUT / "convert_flat.pdf")
    finally:
        pdf.close()
    assert open_document(OUT / "convert_pages.pdf").page_count == 2
    open_document(OUT / "convert_pages.pdf").close()
    assert (OUT / "convert_meta_copy.pdf").is_file()
    assert (OUT / "convert_flat.pdf").is_file()


def test_pdf_to_docx_pptx_xlsx():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    try:
        pdf_to_docx(pdf, OUT / "convert.docx")
        pdf_to_pptx(pdf, OUT / "convert.pptx", pages=[0])
        tables_to_xlsx(pdf, OUT / "convert.xlsx")
    finally:
        pdf.close()
    assert (OUT / "convert.docx").stat().st_size > 100
    assert (OUT / "convert.pptx").stat().st_size > 100
    assert (OUT / "convert.xlsx").stat().st_size > 100


def test_office_to_pdf_unsupported(monkeypatch):
    from core.convert import office_bridge

    monkeypatch.setenv("MONOGRAPH_PDF_DISABLE_OFFICE", "1")
    office_bridge.office_availability(refresh=True)
    try:
        assert office_to_pdf_supported() is False
        with pytest.raises(UnsupportedConversionError):
            office_to_pdf(Path("a.docx"), OUT / "nope.pdf")
    finally:
        monkeypatch.delenv("MONOGRAPH_PDF_DISABLE_OFFICE")
        office_bridge.office_availability(refresh=True)


def test_watermark_header_page_numbers():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    try:
        apply_watermark(pdf, OUT / "convert_wm.pdf", text="TEST WM", opacity=0.3)
        apply_headers_footers(
            pdf, OUT / "convert_hf.pdf", header="Hdr", footer="Ftr"
        )
        apply_page_numbers(pdf, OUT / "convert_pn.pdf")
    finally:
        pdf.close()
    assert (OUT / "convert_wm.pdf").is_file()
    assert (OUT / "convert_hf.pdf").is_file()
    assert (OUT / "convert_pn.pdf").is_file()


def test_image_box_export_and_bake():
    pdf = open_document(SAMPLES / "sample_text.pdf")
    dest = OUT / "convert_image_bake.pdf"
    try:
        edit = EditableDocument(pdf)
        img = ImageBox(
            page_index=0,
            rect=Rect(50, 50, 120, 80),
            source_path=str(IMG_DIR / "swatch.png"),
            opacity=0.9,
            z_index=edit.next_z(),
        )
        edit.add_image(img)
        export_image_box(img, OUT / "convert_exported_img.png")
        assert (OUT / "convert_exported_img.png").is_file()
        bake_into_document(edit, dest)
    finally:
        pdf.close()
    assert dest.is_file()


def test_pdf_to_docx_layout_keeps_tables_and_text(tmp_path):
    pytest.importorskip("pdf2docx")
    from docx import Document as Docx

    from core.convert.pdf_to_docx import layout_engine_available

    assert layout_engine_available()
    src = SAMPLES / "sample_itinerary_table.pdf"
    if not src.is_file():
        from tests.fixtures.make_form_samples import main as make_forms

        make_forms()
    pdf = open_document(src)
    try:
        out = pdf_to_docx(pdf, tmp_path / "layout.docx", engine="layout")
        words = set((pdf.page_text(0) or "").split())
    finally:
        pdf.close()
    doc = Docx(out)
    body = " ".join(p.text for p in doc.paragraphs)
    cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    found = {w for w in words if w in body or w in cells}
    assert len(found) >= 0.8 * len(words), (len(found), len(words))
    assert doc.tables, "table structure should survive"


def test_pdf_to_docx_text_engine_fallback(tmp_path):
    pdf = open_document(SAMPLES / "sample_text.pdf")
    try:
        out = pdf_to_docx(pdf, tmp_path / "plain.docx", engine="text")
    finally:
        pdf.close()
    assert out.stat().st_size > 100


def test_office_to_pdf_round_trip_when_suite_installed(tmp_path):
    from docx import Document as Docx

    if not office_to_pdf_supported(".docx"):
        pytest.skip("no Microsoft Office / LibreOffice on this machine")
    d = Docx()
    d.add_heading("Round trip heading", 1)
    d.add_paragraph("Body text that must survive the conversion.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A1"
    t.cell(1, 1).text = "B2"
    src = tmp_path / "source.docx"
    d.save(src)
    out = office_to_pdf(src, tmp_path / "source.pdf")
    pdf = open_document(out)
    try:
        text = pdf.page_text(0)
    finally:
        pdf.close()
    for needle in ("Round trip heading", "Body text", "A1", "B2"):
        assert needle in text
