"""Area OCR, Extract Text (OCR), Scanned PDF to Word and OCR Languages through the window."""

from __future__ import annotations

import os
from pathlib import Path

import fitz
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("MONOGRAPH_PDF_HEADLESS", "1")
pytest.importorskip("pytestqt")


@pytest.fixture
def win(qtbot, tmp_path, monkeypatch):
    from core import user_prefs
    from core.ocr.paths import ocr_ready
    from ui.main_window import MainWindow, create_application

    ok, detail = ocr_ready()
    if not ok:
        pytest.skip(detail)
    monkeypatch.setattr(user_prefs, "user_data_dir", lambda: tmp_path / "appdata")
    create_application(["monograph-pdf-ocr-ui"])
    src = tmp_path / "scan.pdf"
    page_src = fitz.open()
    p = page_src.new_page(width=600, height=400)
    p.insert_text((40, 80), "TOP LINE OF THE SCAN", fontsize=28)
    p.insert_text((40, 250), "INVOICE NUMBER 12345", fontsize=28)
    pix = p.get_pixmap(dpi=200, alpha=False)
    doc = fitz.open()
    page = doc.new_page(width=600, height=400)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(src)
    w = MainWindow()
    qtbot.addWidget(w)
    w.resize(1200, 800)
    w.show()
    assert w.open_path(src)
    yield w, src
    w.close_document(save_session=False)
    w.close()


def _drag(qtbot, w, page, x0, y0, x1, y1):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    from core.editing.geometry import Point

    inner = w._canvas._canvas  # noqa: SLF001
    p0 = inner.pdf_to_widget(page, Point(x0, y0))
    p1 = inner.pdf_to_widget(page, Point(x1, y1))
    QTest.mousePress(inner, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p0)
    QTest.mouseMove(inner, p1)
    QTest.mouseRelease(inner, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p1)


def _capture_text_dialog(monkeypatch):
    from ui import ocr_dialogs

    shown = []

    def fake_exec(self):
        shown.append(self)
        return 1

    monkeypatch.setattr(ocr_dialogs.OcrTextDialog, "exec", fake_exec)
    return shown


def test_area_ocr_reads_the_dragged_area_and_saves_corrected_text(win, qtbot, monkeypatch, tmp_path):
    w, src = win
    shown = _capture_text_dialog(monkeypatch)
    w.act_ocr_area.trigger()
    _drag(qtbot, w, 0, 20, 210, 580, 280)
    qtbot.waitUntil(lambda: bool(shown), timeout=60000)
    dlg = shown[0]
    text = dlg.text.toPlainText()
    assert "INVOICE" in text and "TOP" not in text
    dlg.text.setPlainText(text.replace("INVOICE", "Invoice"))
    out = dlg.save_text(str(tmp_path / "area.txt"))
    assert "Invoice" in Path(out).read_text(encoding="utf-8-sig")
    docx = dlg.save_docx(str(tmp_path / "area.docx"))
    from docx import Document

    assert any("Invoice" in p.text for p in Document(docx).paragraphs)
    # the open PDF is not changed by OCR text extraction
    assert not w._edit.dirty  # noqa: SLF001
    with fitz.open(src) as d:
        assert d[0].get_text().strip() == ""


def test_extract_text_of_all_pages(win, qtbot, monkeypatch):
    from ui import ocr_dialogs

    w, src = win
    shown = _capture_text_dialog(monkeypatch)

    class Scope(ocr_dialogs.OcrScopeDialog):
        def exec(self):  # noqa: A003
            self.r_all.setChecked(True)
            return self.DialogCode.Accepted

    monkeypatch.setattr(ocr_dialogs, "OcrScopeDialog", Scope)
    w.act_ocr_extract.trigger()
    qtbot.waitUntil(lambda: bool(shown), timeout=60000)
    text = shown[0].text.toPlainText()
    assert "TOP" in text and "INVOICE" in text


def test_scanned_pdf_to_word(win, qtbot, monkeypatch, tmp_path):
    from docx import Document

    from ui import convert_dialogs, ocr_dialogs

    w, src = win
    out = tmp_path / "scan.docx"

    class Scope(ocr_dialogs.OcrScopeDialog):
        def exec(self):  # noqa: A003
            return self.DialogCode.Accepted

    monkeypatch.setattr(ocr_dialogs, "OcrScopeDialog", Scope)
    monkeypatch.setattr(convert_dialogs, "choose_save", lambda *a, **k: out)
    w.act_scan_to_word.trigger()
    qtbot.waitUntil(out.is_file, timeout=60000)
    text = " ".join(p.text for p in Document(str(out)).paragraphs)
    assert "INVOICE" in text


def test_ocr_languages_dialog_saves_the_choice(win, monkeypatch):
    from PySide6.QtCore import Qt

    from ui import ocr_dialogs

    w, src = win

    class Langs(ocr_dialogs.OcrLanguagesDialog):
        def exec(self):  # noqa: A003
            for i in range(self.picker.list.count()):
                item = self.picker.list.item(i)
                item.setCheckState(Qt.CheckState.Checked if item.data(Qt.ItemDataRole.UserRole) == "eng"
                                   else Qt.CheckState.Unchecked)
            assert self.packs.count() >= 7  # optional languages are listed (not installed)
            return self.DialogCode.Accepted

    monkeypatch.setattr(ocr_dialogs, "OcrLanguagesDialog", Langs)
    w.act_ocr_languages.trigger()
    assert w._prefs.ocr_languages == "eng"  # noqa: SLF001
    assert w._ocr_languages_pref() == "eng"  # noqa: SLF001
