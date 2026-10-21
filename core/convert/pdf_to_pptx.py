"""PDF → PPTX with the PDF's own page size.

Slide size
----------
The slide is the PDF page as displayed (``/Rotate`` applied), converted from PDF points to EMU
(1 pt = 12 700 EMU). A page within 2 pt of a standard sheet (A3, A4, A5, Letter, Legal, Tabloid, in
either orientation) gets that sheet's exact size, so an A4 PDF becomes a 210 × 297 mm (portrait) or
297 × 210 mm (landscape) slide — never PowerPoint's 16:9 default.

PowerPoint has one slide size per file. When the pages differ in size or orientation, ``size_policy``
decides (nothing is ever stretched or cropped; every page keeps its proportions):

``true_size``   (default) the slide is the widest × tallest page; every page is centred at its true
                scale (smaller pages get a white margin).
``fit``         the same slide; every page is enlarged or reduced to fill it as far as its proportions allow.
``first_page``  the slide is the first page's size; the other pages are fitted into it.
``custom``      a chosen slide size (``custom_size_pt``: A4, A3, 16:9 widescreen …); every page is fitted.
``split``       one presentation per group of pages of the same size and orientation, each with that
                exact slide size (``pdf_to_pptx_split``).

Sizes outside PowerPoint's limits (1–56 inches per side) are scaled uniformly and reported.

Modes
-----
``layout``    (default) every page is drawn as one picture at the page's true size: the most faithful
              result; the text is not individually editable.
``editable``  the page's graphics, images, table rules and backgrounds are drawn as one picture
              *without its text*, and every text line is added on top as a native text box at the
              same position, with its font name, size, colour, bold / italic and angle. Fonts not
              installed on the PowerPoint computer are substituted by PowerPoint, so line breaks and
              widths can differ slightly; this mode is never pixel-identical.
"""

from __future__ import annotations

import math
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from core.document import PdfDocument

MODES = ("layout", "editable")

LAYOUT_NOTICE = (
    "Layout-preserving: each PDF page is placed as one picture at its true page size. "
    "The text is not editable in PowerPoint."
)
EDITABLE_NOTICE = (
    "Editable text: page graphics are one picture and every text line becomes a PowerPoint text box "
    "at its original position. Fonts that are not installed are substituted by PowerPoint, so line "
    "breaks and widths can differ slightly; the result is not pixel-identical."
)
# Shown before converting (kept as the single "best-effort" statement for both modes).
FIDELITY_NOTICE = (
    "Best-effort: slides use the PDF's page size (A4 stays A4). Layout-preserving mode places each page "
    "as a picture; Editable text mode adds native text boxes over the page graphics."
)

EMU_PER_PT = 12700
EMU_PER_MM = 36000
_MIN_EMU = 914400  # PowerPoint limits a slide side to 1 … 56 inches
_MAX_EMU = 56 * 914400
_SNAP_PT = 2.0

# name, width mm, height mm (portrait)
_STANDARD_SHEETS = (
    ("A3", 297.0, 420.0),
    ("A4", 210.0, 297.0),
    ("A5", 148.0, 210.0),
    ("Letter", 215.9, 279.4),
    ("Legal", 215.9, 355.6),
    ("Tabloid", 279.4, 431.8),
)


@dataclass
class PptxReport:
    mode: str = "layout"
    slide_width_emu: int = 0
    slide_height_emu: int = 0
    slide_label: str = ""
    strategy: str = "uniform"  # "uniform" | "bounding sheet"
    pages: int = 0
    text_boxes: int = 0
    rotated_text: int = 0
    clamped: int = 0
    fonts: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def slide_size_mm(self) -> tuple[float, float]:
        return self.slide_width_emu / EMU_PER_MM, self.slide_height_emu / EMU_PER_MM

    def summary(self) -> str:
        w, h = self.slide_size_mm
        text = f"{self.pages} slide(s), {self.slide_label or f'{w:.0f} × {h:.0f} mm'}"
        if self.strategy != "uniform":
            text += f" ({self.strategy})"
        if self.mode == "editable":
            text += f"; {self.text_boxes} editable text box(es)"
        return text


# ------------------------------------------------------------------------------ sizes
def _snap_page(width_pt: float, height_pt: float) -> tuple[int, int, str]:
    """EMU size of a page, snapped to a standard sheet when within tolerance."""
    for name, wmm, hmm in _STANDARD_SHEETS:
        for portrait in (True, False):
            sw, sh = (wmm, hmm) if portrait else (hmm, wmm)
            if abs(width_pt - sw / 25.4 * 72.0) <= _SNAP_PT and abs(height_pt - sh / 25.4 * 72.0) <= _SNAP_PT:
                return round(sw * EMU_PER_MM), round(sh * EMU_PER_MM), f"{name} {'portrait' if portrait else 'landscape'}"
    return round(width_pt * EMU_PER_PT), round(height_pt * EMU_PER_PT), ""


def _fit_limits(w: int, h: int) -> tuple[int, int, float]:
    """Uniform scale so both sides are inside PowerPoint's limits (1 … 56 in)."""
    scale = 1.0
    longest, shortest = max(w, h), min(w, h)
    if longest > _MAX_EMU:
        scale = _MAX_EMU / longest
    elif shortest < _MIN_EMU:
        scale = _MIN_EMU / shortest
    return round(w * scale), round(h * scale), scale


# ------------------------------------------------------------------------------ fonts
_FAMILY_FIXES = {
    "helvetica": "Arial",
    "arialmt": "Arial",
    "arial": "Arial",
    "times": "Times New Roman",
    "timesnewroman": "Times New Roman",
    "timesnewromanps": "Times New Roman",
    "timesnewromanpsmt": "Times New Roman",
    "courier": "Courier New",
    "couriernew": "Courier New",
    "couriernewps": "Courier New",
    "couriernewpsmt": "Courier New",
    "symbol": "Symbol",
    "zapfdingbats": "Wingdings",
}


def powerpoint_family(pdf_font: str) -> str:
    """A family name PowerPoint understands for a PDF font name (subset prefix and style removed)."""
    name = pdf_font or ""
    if "+" in name[:8]:
        name = name.split("+", 1)[1]
    base = re.split(r"[-,]", name, maxsplit=1)[0].strip()
    key = re.sub(r"[\s_]", "", base).lower()
    if key in _FAMILY_FIXES:
        return _FAMILY_FIXES[key]
    for suffix in ("psmt", "mt", "ps"):
        if key.endswith(suffix) and key[: -len(suffix)] in _FAMILY_FIXES:
            return _FAMILY_FIXES[key[: -len(suffix)]]
    spaced = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", base).strip()
    return spaced or "Arial"


# ------------------------------------------------------------------------------ text
def _display_lines(page: fitz.Page) -> list[dict[str, Any]]:
    """Text lines as display-space rectangles with angle and styled runs."""
    rot = page.rotation_matrix
    linear = fitz.Matrix(rot.a, rot.b, rot.c, rot.d, 0, 0)
    out: list[dict[str, Any]] = []
    data = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_PRESERVE_LIGATURES) or {}
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = [s for s in line.get("spans", []) if (s.get("text") or "").strip() or s.get("text")]
            if not spans or not "".join(s["text"] for s in spans).strip():
                continue
            rect = fitz.Rect(line["bbox"]) * rot
            rect.normalize()
            dx, dy = line.get("dir", (1.0, 0.0))
            v = fitz.Point(dx, dy) * linear
            angle = math.degrees(math.atan2(v.y, v.x))
            out.append({"rect": rect, "angle": angle, "spans": spans, "origin": fitz.Point(spans[0]["origin"]) * rot})
    return out


def _add_text_line(slide: Any, line: dict[str, Any], sx: float, sy: float, ox: int, oy: int,
                   slide_w: int, slide_h: int, report: PptxReport) -> None:
    rect: fitz.Rect = line["rect"]
    angle = ((line["angle"] + 180.0) % 360.0) - 180.0
    rotated = abs(angle) > 0.5
    if rotated:
        # the line's own (unrotated) extent: length along the baseline × text height
        length = max(rect.width, rect.height) if abs(abs(angle) - 90.0) < 5.0 or abs(angle) > 175.0 else math.hypot(rect.width, rect.height)
        height = max(s["size"] for s in line["spans"]) * 1.25
        cx, cy = (rect.x0 + rect.x1) / 2.0, (rect.y0 + rect.y1) / 2.0
        w_emu, h_emu = round(length * sx), round(height * sy)
        ccx, ccy = round(cx * sx) + ox, round(cy * sy) + oy
        # keep the *turned* box (its visual extent, not the unturned frame) inside the slide
        rad = math.radians(angle)
        half_w = (abs(w_emu * math.cos(rad)) + abs(h_emu * math.sin(rad))) / 2.0
        half_h = (abs(w_emu * math.sin(rad)) + abs(h_emu * math.cos(rad))) / 2.0
        nx = min(max(ccx, half_w), max(half_w, slide_w - half_w))
        ny = min(max(ccy, half_h), max(half_h, slide_h - half_h))
        if (nx, ny) != (ccx, ccy):
            report.clamped += 1
        ccx, ccy = round(nx), round(ny)
        left, top = ccx - w_emu // 2, ccy - h_emu // 2
        report.rotated_text += 1
    else:
        left, top = round(rect.x0 * sx) + ox, round(rect.y0 * sy) + oy
        w_emu, h_emu = round(rect.width * sx), round(rect.height * sy)
    w_emu, h_emu = max(w_emu, 12700), max(h_emu, 12700)
    if not rotated:  # never place a box (or the text in it) outside the slide
        clamped_left = min(max(0, left), max(0, slide_w - w_emu))
        clamped_top = min(max(0, top), max(0, slide_h - h_emu))
        if (clamped_left, clamped_top) != (left, top) or w_emu > slide_w or h_emu > slide_h:
            report.clamped += 1
        left, top = clamped_left, clamped_top
        w_emu, h_emu = min(w_emu, slide_w), min(h_emu, slide_h)
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(w_emu), Emu(h_emu))
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    para = frame.paragraphs[0]
    for span in line["spans"]:
        text = span.get("text") or ""
        if not text:
            continue
        run = para.add_run()
        run.text = text
        font = run.font
        family = powerpoint_family(str(span.get("font", "")))
        font.name = family
        if family not in report.fonts:
            report.fonts.append(family)
        font.size = Pt(max(1.0, round(float(span.get("size", 12.0)) * sy / EMU_PER_PT, 2)))
        flags = int(span.get("flags", 0))
        low = str(span.get("font", "")).lower()
        font.bold = bool(flags & 16) or "bold" in low or "black" in low
        font.italic = bool(flags & 2) or "italic" in low or "oblique" in low
        if flags & 1:  # superscript
            run._r.get_or_add_rPr().set("baseline", "30000")
        color = int(span.get("color", 0))
        font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
    if rotated:
        box.rotation = angle % 360.0
    report.text_boxes += 1


# ------------------------------------------------------------------------------ conversion
def _live_bytes(document: PdfDocument) -> bytes:
    doc: fitz.Document = document._backend._doc  # noqa: SLF001
    return doc.tobytes(garbage=1, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)


def _render(page: fitz.Page, dpi: int, path: Path) -> None:
    long_edge_pt = max(page.rect.width, page.rect.height)
    zoom = min(dpi / 72.0, 6000.0 / max(long_edge_pt, 1.0))
    page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).save(path.as_posix())


SIZE_POLICIES = ("true_size", "fit", "first_page", "custom", "split")
SLIDE_PRESETS_PT = {
    "A4 portrait": (595.28, 841.89), "A4 landscape": (841.89, 595.28),
    "A3 portrait": (841.89, 1190.55), "A3 landscape": (1190.55, 841.89),
    "Letter portrait": (612.0, 792.0), "Letter landscape": (792.0, 612.0),
    "Widescreen 16:9": (960.0, 540.0), "Standard 4:3": (720.0, 540.0),
}


def page_groups(document: PdfDocument, pages: list[int] | None = None) -> list[tuple[str, list[int]]]:
    """Pages grouped by size and orientation (2 pt tolerance), in order of first appearance:
    [("A4 portrait", [0, 2]), ("A3 landscape", [1]), ...]."""
    doc = document._backend._doc
    indices = list(range(doc.page_count)) if pages is None else list(pages)
    groups: list[tuple[tuple[int, int, str], list[int]]] = []
    for i in indices:
        r = doc.load_page(i).rect
        ew, eh, label = _snap_page(float(r.width), float(r.height))
        for key, members in groups:
            if abs(key[0] - ew) <= 2 * EMU_PER_PT and abs(key[1] - eh) <= 2 * EMU_PER_PT:
                members.append(i)
                break
        else:
            name = label or f"{r.width / 72 * 25.4:.0f} × {r.height / 72 * 25.4:.0f} mm"
            groups.append(((ew, eh, name), [i]))
    return [(key[2], members) for key, members in groups]


def pdf_to_pptx_split(document: PdfDocument, destination: Path, *, pages: list[int] | None = None,
                      **kwargs) -> list[Path]:
    """One presentation per page size/orientation group: ``name (A4 portrait).pptx`` …; a document with
    one size gives one file at ``destination``."""
    destination = Path(destination)
    groups = page_groups(document, pages)
    if len(groups) <= 1:
        return [pdf_to_pptx(document, destination, pages=pages, **kwargs)]
    out = []
    for label, members in groups:
        safe = re.sub(r'[\\/:*?"<>|]', "-", label)
        target = destination.with_name(f"{destination.stem} ({safe}){destination.suffix}")
        out.append(pdf_to_pptx(document, target, pages=members, **kwargs))
    return out


def pdf_to_pptx(
    document: PdfDocument,
    destination: Path,
    *,
    pages: list[int] | None = None,
    dpi: int = 200,
    mode: str = "layout",
    report: PptxReport | None = None,
    size_policy: str = "true_size",
    custom_size_pt: tuple[float, float] | None = None,
) -> Path:
    """Write ``destination`` (.pptx). ``report`` (optional) receives size, strategy and limitations."""
    if size_policy not in SIZE_POLICIES or size_policy == "split":
        raise ValueError(f"Unsupported slide size policy here: {size_policy} (use pdf_to_pptx_split for split)")
    if size_policy == "custom" and not custom_size_pt:
        raise ValueError("custom slide size requires custom_size_pt")
    if mode not in MODES:
        raise ValueError(f"Unsupported PowerPoint mode: {mode} (supported: {', '.join(MODES)})")
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    report = report if report is not None else PptxReport()
    report.mode = mode

    src = fitz.open("pdf", _live_bytes(document))
    try:
        indices = list(range(src.page_count)) if pages is None else list(pages)
        for i in indices:
            if i < 0 or i >= src.page_count:
                raise IndexError(f"page out of range: {i}")
        if not indices:
            raise ValueError("No pages to convert")

        sizes = []  # per page: (emu_w, emu_h, label, pt_w, pt_h)
        for i in indices:
            r = src.load_page(i).rect  # already has /Rotate applied
            ew, eh, label = _snap_page(float(r.width), float(r.height))
            sizes.append((ew, eh, label, float(r.width), float(r.height)))

        uniform = all(abs(s[0] - sizes[0][0]) <= 2 * EMU_PER_PT and abs(s[1] - sizes[0][1]) <= 2 * EMU_PER_PT for s in sizes)
        fit_pages = False
        if size_policy == "custom":
            slide_w, slide_h, label = _snap_page(*custom_size_pt)
            label = label or f"{custom_size_pt[0] / 72 * 25.4:.0f} × {custom_size_pt[1] / 72 * 25.4:.0f} mm"
            report.strategy, fit_pages = "custom slide size", True
            report.notes.append("Every page is fitted into the chosen slide size, keeping its proportions.")
        elif uniform:
            slide_w, slide_h, label = sizes[0][0], sizes[0][1], sizes[0][2]
            report.strategy = "uniform"
        elif size_policy == "first_page":
            slide_w, slide_h, label = sizes[0][0], sizes[0][1], sizes[0][2]
            report.strategy, fit_pages = "first page size", True
            report.notes.append("The pages differ in size: the slide is the first page's size and the other "
                                "pages are fitted into it, keeping their proportions.")
        else:
            slide_w, slide_h = max(s[0] for s in sizes), max(s[1] for s in sizes)
            label = ""
            fit_pages = size_policy == "fit"
            report.strategy = "bounding sheet, pages fitted" if fit_pages else "bounding sheet"
            report.notes.append(
                "The pages differ in size or orientation. PowerPoint allows one slide size, so the slide is the "
                "widest × tallest page and every page is "
                + ("enlarged or reduced to fill it, keeping its proportions." if fit_pages else
                   "centered at its true size, without stretching or cropping.")
            )
        slide_w, slide_h, k = _fit_limits(slide_w, slide_h)
        if k != 1.0:
            report.notes.append("The page is outside PowerPoint's slide size limits (1–56 in) and was scaled uniformly.")
            label = ""
        report.slide_width_emu, report.slide_height_emu, report.slide_label = slide_w, slide_h, label
        report.pages = len(indices)

        prs = Presentation()
        prs.slide_width = Emu(slide_w)
        prs.slide_height = Emu(slide_h)
        blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]

        with tempfile.TemporaryDirectory(prefix="monographpdf_pptx_") as tmp:
            tmpdir = Path(tmp)
            for n, i in enumerate(indices):
                ew, eh, _lab, pw, ph = sizes[n]
                if fit_pages:
                    f = min(slide_w / ew, slide_h / eh)
                    ew, eh = round(ew * f), round(eh * f)
                else:
                    ew, eh = round(ew * k), round(eh * k)
                ox, oy = (slide_w - ew) // 2, (slide_h - eh) // 2
                sx, sy = ew / max(pw, 1.0), eh / max(ph, 1.0)
                page = src.load_page(i)
                lines = _display_lines(page) if mode == "editable" else []
                if mode == "editable":
                    for block in page.get_text("blocks"):
                        if len(block) > 6 and block[6] == 0:
                            page.add_redact_annot(fitz.Rect(block[:4]))
                    page.apply_redactions(
                        images=fitz.PDF_REDACT_IMAGE_NONE,
                        graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                        text=fitz.PDF_REDACT_TEXT_REMOVE,
                    )
                png = tmpdir / f"page_{n:04d}.png"
                _render(page, dpi, png)
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(str(png), Emu(ox), Emu(oy), width=Emu(ew), height=Emu(eh))
                for line in lines:
                    _add_text_line(slide, line, sx, sy, ox, oy, slide_w, slide_h, report)
            if mode == "editable":
                report.notes.append(
                    "Editable text: fonts not installed on the PowerPoint computer are substituted, so line breaks "
                    "and widths can differ slightly; the page is not pixel-identical."
                )
            else:
                report.notes.append("Layout-preserving: pages are pictures; text is not editable.")
            from core.convert.office_metadata import (
                apply_core_properties,
                pdf_properties,
            )

            apply_core_properties(prs.core_properties, pdf_properties(src.metadata))
            if report.fonts:
                report.notes.append("Fonts used for text boxes: " + ", ".join(report.fonts[:12])
                                    + ". PowerPoint substitutes any that are not installed.")
            tmp_out = destination.with_name(destination.name + ".part")
            prs.save(tmp_out.as_posix())
            tmp_out.replace(destination)
    finally:
        src.close()
    return destination
