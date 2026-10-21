"""PDF → Word (DOCX).

Engines
-------
``layout`` (default)  pdf2docx — rebuilds each page as editable Word content and keeps
                      the original layout: fonts, sizes, colours, paragraphs, tables,
                      images, columns and page size. Fully offline, bundled.
``word``              Microsoft Word's own PDF Reflow (only when Word is installed) —
                      best for long flowing text you want to re-edit heavily.
``text``              Plain paragraphs only (legacy fallback; always available).
"""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Callable

from core.document import PdfDocument

FIDELITY_NOTICE = (
    "Layout-preserving conversion: text, fonts, colors, tables, images and page size are "
    "rebuilt as editable Word content. Very complex pages (overlapping text boxes, unusual "
    "fonts, scanned images without OCR) may need small touch-ups in Word."
)

ENGINES = ("layout", "word", "text")


def layout_engine_available() -> bool:
    try:
        import pdf2docx  # noqa: F401

        return True
    except Exception:  # noqa: BLE001
        return False


def word_engine_available() -> bool:
    from core.convert.office_bridge import office_availability

    return office_availability().word


def _live_pdf_bytes(document: PdfDocument) -> bytes:
    """Current in-memory state (page edits, rotations, annotations), decrypted."""
    import fitz

    doc: fitz.Document = document._backend._doc  # noqa: SLF001
    return doc.tobytes(garbage=1, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)


def _page_sizes_pt(document: PdfDocument, indices: list[int]) -> list[tuple[float, float]]:
    doc = document._backend._doc  # noqa: SLF001
    return [(float(doc.load_page(i).rect.width), float(doc.load_page(i).rect.height)) for i in indices]


def _set_section_geometry(section, width_pt: float, height_pt: float, *, margins: bool = True) -> None:
    """Word section = the PDF page: exact size, orientation flag matching the shape, and (for the text
    mode) margins of 0.75 in, at most 10 % of the page side."""
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    section.orientation = WD_ORIENT.LANDSCAPE if width_pt > height_pt else WD_ORIENT.PORTRAIT
    section.page_width = Pt(width_pt)
    section.page_height = Pt(height_pt)
    if margins:
        mx, my = Pt(min(54.0, width_pt * 0.1)), Pt(min(54.0, height_pt * 0.1))
        section.left_margin = section.right_margin = mx
        section.top_margin = section.bottom_margin = my


def _pdf_to_docx_text(document: PdfDocument, destination: Path, pages: list[int] | None) -> Path:
    """Editable text: one Word section per PDF page with that page's size and orientation (mixed sizes
    stay mixed); the text flows as plain paragraphs inside it."""
    from docx import Document as DocxDocument
    from docx.enum.section import WD_SECTION

    docx = DocxDocument()
    indices = list(range(document.page_count)) if pages is None else list(pages)
    for pi in indices:
        if pi < 0 or pi >= document.page_count:
            raise IndexError(f"page out of range: {pi}")
    sizes = _page_sizes_pt(document, indices)
    for i, pi in enumerate(indices):
        section = docx.sections[0] if i == 0 else docx.add_section(WD_SECTION.NEW_PAGE)
        _set_section_geometry(section, *sizes[i])
        text = document.page_text(pi) or ""
        if not text.strip():
            docx.add_paragraph("[No extractable text on this page — run OCR first]")
        else:
            for block in text.split("\n"):
                docx.add_paragraph(block)
    tmp = destination.with_name(destination.name + ".part")
    docx.save(tmp.as_posix())
    tmp.replace(destination)
    return destination


def _fix_orientation(path: Path) -> None:
    """pdf2docx sets each section's exact size but always marks it portrait; Word and printers read the
    orientation flag, so landscape-shaped sections are marked landscape (size unchanged)."""
    from docx import Document as DocxDocument
    from docx.enum.section import WD_ORIENT

    docx = DocxDocument(path.as_posix())
    changed = False
    for section in docx.sections:
        landscape = section.page_width and section.page_height and section.page_width > section.page_height
        want = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if section.orientation != want:
            section.orientation = want
            changed = True
    if changed:
        tmp = path.with_name(path.name + ".orient")
        docx.save(tmp.as_posix())
        tmp.replace(path)


def _pdf_to_docx_layout(
    document: PdfDocument,
    destination: Path,
    pages: list[int] | None,
    progress: Callable[[int, int], None] | None,
) -> Path:
    import logging

    from pdf2docx import Converter

    logging.getLogger("root").setLevel(logging.WARNING)
    data = _live_pdf_bytes(document)
    cv = Converter(stream=data)
    try:
        kwargs = {"multi_processing": False}
        if pages is not None:
            kwargs["pages"] = [int(p) for p in pages]
        if progress is not None:
            progress(0, 1)
        tmp = destination.with_name(destination.name + ".part")
        try:
            cv.convert(tmp.as_posix(), **kwargs)
            tmp.replace(destination)
        finally:
            tmp.unlink(missing_ok=True)
        if progress is not None:
            progress(1, 1)
    finally:
        cv.close()
    return destination


def _pdf_to_docx_word(document: PdfDocument, destination: Path, pages: list[int] | None) -> Path:
    import fitz

    from core.convert.office_bridge import convert_pdf_to_docx_with_word

    with tempfile.TemporaryDirectory(prefix="monographpdf-word-") as tmp:
        src = Path(tmp) / "source.pdf"
        data = _live_pdf_bytes(document)
        if pages is not None:
            full = fitz.open(stream=data, filetype="pdf")
            part = fitz.open()
            for p in pages:
                part.insert_pdf(full, from_page=int(p), to_page=int(p))
            data = part.tobytes(garbage=1, deflate=True)
            part.close()
            full.close()
        src.write_bytes(data)
        return convert_pdf_to_docx_with_word(src, destination)


def pdf_to_docx(
    document: PdfDocument,
    destination: Path,
    *,
    pages: list[int] | None = None,
    engine: str = "layout",
    progress: Callable[[int, int], None] | None = None,
) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if engine not in ENGINES:
        raise ValueError(f"unknown engine {engine!r}; choose one of {ENGINES}")
    if engine == "layout" and not layout_engine_available():
        engine = "text"
    if engine == "word":
        out = _pdf_to_docx_word(document, destination, pages)
    elif engine == "layout":
        out = _pdf_to_docx_layout(document, destination, pages, progress)
        _fix_orientation(out)
    else:
        out = _pdf_to_docx_text(document, destination, pages)
    # replace the template/converter properties (and Word's user name) with the PDF's own
    from core.convert.office_metadata import apply_to_file, document_metadata

    apply_to_file(out, document_metadata(document))
    return out
